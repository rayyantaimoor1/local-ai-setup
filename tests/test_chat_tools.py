"""Chat with tools: documents, multi-round tool use, git and the push dialog, without any agent.

A scripted fake model stands in for Ollama. Git tests use real repositories with a bare remote.
"""
import json
import subprocess
import sys
import threading
import time
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import approvals as approvals_mod  # noqa: E402
import config  # noqa: E402
import constitution as con  # noqa: E402
import documents  # noqa: E402
import ollama_client  # noqa: E402
import router_engine  # noqa: E402
import tools  # noqa: E402
from approvals import approvals  # noqa: E402

AWS = "AKIA" + "ABCDEFGHIJKLMNOP"
CODER = config.SPECIALISTS["CODE_TASK"]


@pytest.fixture(autouse=True)
def clean(tmp_path):
    """Fresh outputs folder, settings and audit log for every test."""
    for p in (config.CONSTITUTION_SETTINGS, config.AUDIT_LOG):
        Path(p).unlink(missing_ok=True)
    for f in documents.outputs_dir().iterdir():
        f.unlink()
    con.overrides.revoke()
    yield


@pytest.fixture
def model(monkeypatch):
    """Scripted model: `replies` are returned in order (router calls are answered separately)."""
    state = {"replies": [], "prompts": [], "models": [], "keep_alive": [], "unloaded": [], "router": "GENERAL_TASK",
             "think": [], "router_calls": 0}

    def stream(model_name, prompt, system=None, num_ctx=0, timeout=0, temperature=None, cancel=None, keep_alive=None,
               think=None):
        state["prompts"].append(prompt)
        state["models"].append(model_name)
        state["keep_alive"].append(keep_alive)
        state["think"].append(think)
        state["system"] = system
        text = state["replies"].pop(0)
        for i in range(0, len(text), 9):
            yield ("response", text[i:i + 9])

    monkeypatch.setattr(ollama_client, "stream_generate", stream)
    def router(*a, **k):
        state["router_calls"] += 1
        return state["router"]
    monkeypatch.setattr(ollama_client, "generate", router)
    monkeypatch.setattr(ollama_client, "unload", lambda m: state["unloaded"].append(m))
    monkeypatch.setattr(ollama_client, "list_models", lambda: [])
    monkeypatch.setattr(router_engine, "_trim", lambda: {"freed_mb": 0})
    return state


def chat(message, deep=False, **kw):
    events = list(router_engine.process_request(message, "", deep, **kw))
    return events, events[-1]


# --- documents -----------------------------------------------------------------------------------------

def test_parse_markdown_sections_bullets_tables_images():
    md = ("Intro paragraph line one\nline two.\n\nSecond paragraph.\n\n## Points\n- one\n- **two**\n\n"
          "## Table\n| Item | Cost |\n|---|---|\n| Pens | 12 |\n\n## Chart\n![Sales by month](sales.png)\n\n# Last\n1. first\n2. second")
    sections = documents.parse_markdown(md)
    assert sections[0] == ("", [{"text": "Intro paragraph line one line two."}, {"text": "Second paragraph."}])
    assert sections[1] == ("Points", [{"bullets": ["one", "two"]}])
    assert sections[2] == ("Table", [{"table": [["Item", "Cost"], ["Pens", "12"]]}])
    assert sections[3][0] == "Chart" and sections[3][1][0]["caption"] == "Sales by month"
    assert Path(sections[3][1][0]["image"]).parent == documents.outputs_dir()  # images only come from outputs/
    assert sections[4] == ("Last", [{"bullets": ["first", "second"]}])


def test_mixed_section_keeps_bullets_and_table_apart():
    """Regression: a bullet list followed by a table used to be read as one big table."""
    out = documents.create_document("mix.docx", "Mix", "## Both\n- a\n- b\n\n| x | y |\n| 1 | 2 |\n\nAfter text.")
    text = tools.read_docx_text(str(documents.outputs_dir() / "mix.docx"))
    assert out.startswith("Created") and "x | y" in text and "1 | 2" in text and "a\nb" in text and "After text." in text


@pytest.mark.parametrize("filename,check", [
    ("report.docx", lambda p: "Pens | 12" in tools.read_docx_text(str(p)) and "Summary" in tools.read_docx_text(str(p))),
    ("report.pdf", lambda p: "Pens" in __import__("pymupdf").open(p)[0].get_text()),
    ("deck.pptx", lambda p: len(__import__("pptx").Presentation(p).slides) == 3),
    ("data.xlsx", lambda p: __import__("openpyxl").load_workbook(p)["Costs"]["B2"].value == 12),
    ("data.csv", lambda p: p.read_text(encoding="utf-8-sig").splitlines()[1] == "Pens,12"),
    ("notes.md", lambda p: p.read_text(encoding="utf-8").startswith("# Report")),
])
def test_create_document_every_format(filename, check):
    md = "## Summary\nAll good.\n\n## Costs\n| Item | Cost |\n|---|---|\n| Pens | 12 |"
    out = documents.create_document(filename, "Report", md)
    assert out.startswith("Created "), out
    path = documents.outputs_dir() / filename
    assert path.exists() and check(path)


def test_documents_never_overwrite_and_never_escape_the_outputs_folder():
    assert documents.create_document("r.docx", "A", "x").startswith("Created r.docx")
    assert documents.create_document("r.docx", "B", "y").startswith("Created r-2.docx")
    assert documents.create_document("..\\..\\evil.docx", "C", "z").startswith("Created evil.docx")
    assert documents.create_document("C:\\Windows\\x.pdf", "D", "z").startswith("Created x.pdf")
    assert not (ROOT / "evil.docx").exists() and not (ROOT.parent / "evil.docx").exists()
    assert documents.create_document("bad.exe", "E", "z").startswith("Error: Unsupported file type")
    assert documents.create_document("noext", "F", "z").startswith("Created noext.docx")


def test_spreadsheet_needs_data_and_accepts_csv_lines():
    assert "Markdown table" in documents.create_document("x.xlsx", "T", "just prose")
    assert documents.create_document("y.xlsx", "T", "n,square\n1,1\n2,4").startswith("Created")
    import openpyxl
    ws = openpyxl.load_workbook(documents.outputs_dir() / "y.xlsx").active
    assert [c.value for c in ws[3]] == [2, 4]


def test_secrets_are_never_copied_into_documents():
    out = documents.create_document("keys.docx", "Keys", f"Our AWS key is {AWS}")
    assert out.startswith("Error: Refused by the constitution") and not (documents.outputs_dir() / "keys.docx").exists()
    assert any(e["event"] == "document_blocked" and e["rule"] == "secrets_in_documents" for e in con.read_audit())
    assert "secrets_in_documents" in {r["id"] for r in con.rules_view() if r["locked"]}


def test_chart_and_document_with_chart():
    assert documents.create_chart("sales.png", "bar", ["Jan", "Feb"], [10, 12], "Sales").startswith("Created sales.png")
    assert documents.create_chart("x.png", "radar", ["a"], [1]).startswith("Error: kind")
    assert documents.create_chart("y.png", "bar", ["a", "b"], [1]).startswith("Error: labels")
    out = documents.create_document("r.docx", "R", "## Chart\n![Sales](sales.png)")
    from docx import Document
    assert out.startswith("Created") and len(Document(documents.outputs_dir() / "r.docx").inline_shapes) == 1


def test_resolve_output_is_confined():
    for bad in ("../x", "..\\x", "C:\\Windows\\win.ini", "/etc/passwd", ""):
        with pytest.raises(ValueError):
            documents.resolve_output(bad)


# --- parsing tool calls ------------------------------------------------------------------------------------

def test_multi_line_tool_call_with_triple_quoted_content():
    text = ('TOOL_CALL: create_document(filename="r.docx", title="Report (draft)", content="""## A\n'
            'Costs (approx): see table\n| a | b |\n| 1 | 2 |""")')
    [(name, args, kwargs)] = tools.parse_tool_calls(text)
    assert name == "create_document" and kwargs["title"] == "Report (draft)" and "| 1 | 2 |" in kwargs["content"]


def test_several_calls_and_junk():
    text = 'TOOL_CALL: git_status()\nTOOL_CALL: broken(\nTOOL_CALL: git_log(n=3)\nTOOL_CALL: __import__("os").system("x")'
    assert [c[0] for c in tools.parse_tool_calls(text)] == ["git_status", "git_log"]


@pytest.mark.parametrize("typed,real", [
    ("make_docx", "create_document"), ("create_report", "create_document"), ("make_chart", "create_chart"),
    ("git_push", "git_push_request"), ("read_pdf_text", "extract_pdf_text"), ("list_files", "get_directory_tree"),
    ("creat_document", "create_document"),
])
def test_made_up_tool_names_map_to_real_ones(typed, real):
    assert tools.resolve_tool_name(typed, tools.MODEL_TOOLS) == real


def test_no_tool_can_run_code_delete_or_change_the_rules():
    for name in tools.MODEL_TOOLS:
        assert not any(bad in name for bad in ("run_python", "exec", "delete", "remove", "pip", "shell",
                                               "override", "setting", "constitution", "write_file"))


# --- the chat tool loop -------------------------------------------------------------------------------------

def test_pdf_to_word_report_in_two_model_calls(model, tmp_path):
    import pymupdf
    pdf = tmp_path / "invoice.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "INVOICE 4471. Customer: Acme Ltd. Total due: 1,280 USD.")
    doc.save(pdf)
    doc.close()
    model["replies"] += [
        f"TOOL_CALL: extract_pdf_text(path={str(pdf)!r})",
        'TOOL_CALL: create_document(filename="invoice_report.docx", title="Invoice", content="""## Summary\n'
        '- Customer: Acme Ltd.\n- Total due: 1,280 USD""")',
    ]
    events, done = chat(f"Create a Word report summarizing the PDF at {pdf}")
    assert model["models"] == [CODER, CODER]  # no router call (rule), exactly two generations
    assert "Acme Ltd." in model["prompts"][1] and "untrusted data" in model["prompts"][1]
    assert [e["name"] for e in events if e["type"] == "tool"] == ["extract_pdf_text", "create_document"]
    assert done["reply"].startswith("Created invoice_report.docx")
    assert done["meta"]["files"] == ["invoice_report.docx"] and done["meta"]["rounds"] == 2
    assert "Acme Ltd." in tools.read_docx_text(str(documents.outputs_dir() / "invoice_report.docx"))
    assert set(model["keep_alive"]) == {config.IDLE_KEEP_ALIVE} and model["unloaded"] == []  # kept for a follow-up


def test_read_then_answer(model, tmp_path):
    (tmp_path / "a.txt").write_text("The answer is 42.")
    model["replies"] += [f"TOOL_CALL: read_file(path={str(tmp_path / 'a.txt')!r})", "It says the answer is 42."]
    events, done = chat(f"What does the file {tmp_path / 'a.txt'} say? Show me the files")
    assert done["reply"] == "It says the answer is 42." and done["meta"]["tool"] == "read_file"
    assert "TOOL_CALL" not in "".join(e["text"] for e in events if e["type"] == "token")


def test_tool_error_lets_the_model_fix_it(model):
    model["replies"] += ['TOOL_CALL: create_document(filename="d.xlsx", title="T", content="no table here")',
                        'TOOL_CALL: create_document(filename="d.xlsx", title="T", content="| n | sq |\\n| 2 | 4 |")']
    events, done = chat("Make an Excel spreadsheet of squares")
    assert done["reply"].startswith("Created d.xlsx") and done["meta"]["rounds"] == 2
    assert "Markdown table" in model["prompts"][1]  # the error was shown to the model


def test_tool_rounds_are_capped_and_the_last_round_must_answer(model):
    calls = ["TOOL_CALL: get_system_info()", 'TOOL_CALL: get_directory_tree(path=".")', 'TOOL_CALL: git_status()']
    model["replies"] += calls[:config.CHAT_MAX_TOOL_ROUNDS - 1] + ["Here is my answer."]
    events, done = chat("How much free RAM does this computer have right now?")
    assert done["reply"] == "Here is my answer." and done["meta"]["rounds"] == config.CHAT_MAX_TOOL_ROUNDS
    assert "No more tools can be used" in model["prompts"][-1]


def test_a_model_that_only_wants_more_tools_gets_an_honest_fallback(model):
    calls = ["TOOL_CALL: get_system_info()", 'TOOL_CALL: get_directory_tree(path=".")', 'TOOL_CALL: git_status()']
    model["replies"] += calls[:config.CHAT_MAX_TOOL_ROUNDS - 1] + ['TOOL_CALL: search_files(pattern="x")']
    events, done = chat("How much free RAM does this computer have right now?")
    assert done["reply"].startswith("I ran out of tool steps") and "(The model returned no answer.)" not in done["reply"]
    assert sum(e["type"] == "tool" for e in events) == config.CHAT_MAX_TOOL_ROUNDS - 1  # the last request did not run


def test_plain_chat_is_one_call_with_no_tools(model):
    model["router"] = "GENERAL_TASK"
    model["replies"] += ["Hello! TOOL_CALL: get_system_info()"]  # not a tool category: nothing runs
    events, done = chat("hi there")
    assert model["models"] == [config.SPECIALISTS["GENERAL_TASK"]] and model["keep_alive"] == [config.IDLE_KEEP_ALIVE]
    assert not any(e["type"] == "tool" for e in events) and model["unloaded"] == []
    assert "TOOL_CALL" not in model["system"]
    assert "CONSTITUTION" not in model["system"] and "constitution" not in model["system"].lower()  # enforced in code only


def test_document_with_a_secret_is_refused_in_chat(model):
    model["replies"] += [f'TOOL_CALL: create_document(filename="k.docx", title="K", content="key {AWS}")',
                         "I could not create it: the content looked like a secret."]
    events, done = chat("Create a Word document with my key")
    assert "secret" in done["reply"] and done["meta"]["files"] == []


def test_early_exit_waits_until_every_requested_file_exists(model):
    """Real failure: 'an Excel sheet AND a bar chart' stopped after the spreadsheet."""
    model["replies"] += [
        'TOOL_CALL: create_document(filename="sq.xlsx", title="Squares", content="| n | sq |\\n| 1 | 1 |\\n| 2 | 4 |")',
        'TOOL_CALL: create_chart(filename="sq.png", kind="bar", labels=["1", "2"], values=[1, 4], title="Squares")',
    ]
    events, done = chat("Make an Excel spreadsheet of the first 2 square numbers and a bar chart of them")
    assert done["meta"]["files"] == ["sq.xlsx", "sq.png"] and done["meta"]["rounds"] == 2
    assert "also asks for .png" in model["prompts"][1]
    assert done["reply"].startswith("Created sq.png")


def test_a_single_requested_file_still_ends_at_once(model):
    model["replies"] += ['TOOL_CALL: create_document(filename="r.docx", title="R", content="hi")']
    events, done = chat("Create a Word report saying hi")
    assert done["meta"]["rounds"] == 1


@pytest.mark.parametrize("request_text,exts", [
    ("Create a Word report about cats", [".docx"]),
    ("Make an Excel spreadsheet and a chart of sales", [".xlsx", ".png"]),
    ("Build a PowerPoint presentation on Python", [".pptx"]),
    ("Generate a PDF of the results", [".pdf"]),
    ("Export the data as csv", [".csv"]),
    ("Extract the text from invoice.pdf", []),
    ("Summarize the report in one word", []),
    (r"Create a Word report summarizing the PDF at C:\x\invoice.pdf. Include the total.", [".docx"]),
    ("Make an Excel spreadsheet of the first 10 square numbers (columns n and n squared) and a bar chart of them.", [".xlsx", ".png"]),
    ("Read the PowerPoint presentation deck.pptx and list its titles", []),
    ("Search my project for TODO comments and put them in a Word table", [".docx"]),
])
def test_requested_outputs(request_text, exts):
    assert documents.requested_outputs(request_text) == exts


# --- deep thinking on / off ------------------------------------------------------------------------------------

def test_math_uses_qwen3_in_one_pass_when_fast(model):
    """Benchmark winner: qwen3:8b without its reasoning phase (2/3 right vs 1/3 for llama3.1 and deepseek-r1)."""
    model["router"] = "MATH_TASK"
    model["replies"] += ["447"]
    events, done = chat("What is 15% of 2400 plus 87?")
    assert model["models"] == ["qwen3:8b"] and done["meta"]["thinking"] == "fast" and done["reply"] == "447"
    assert "never LaTeX" in model["system"]


def test_deep_math_runs_the_refine_loop_with_qwen3(model):
    model["router"] = "MATH_TASK"
    model["replies"] += ["35 min = 0.75 h, so 76.5 km/h", "- 35 minutes is 0.583 h, not 0.75 h", "80.1 km/h", "NO ISSUES"]
    events, done = chat("A train covers 287 km from 9:40 to 13:15. Average speed?", think=True)
    assert model["models"] == ["qwen3:8b"] * 4 and done["reply"] == "80.1 km/h"
    assert "0.583 h" in model["prompts"][2]


def test_a_reasoning_model_can_still_be_configured(model, monkeypatch):
    """Putting deepseek-r1 back (THINKING_MODELS + a fast stand-in) restores native reasoning in deep mode."""
    monkeypatch.setattr(config, "SPECIALISTS", {**config.SPECIALISTS, "MATH_TASK": "deepseek-r1:8b"})
    monkeypatch.setattr(config, "THINKING_MODELS", ("deepseek-r1:8b",))
    monkeypatch.setattr(config, "FAST_MODELS", {"MATH_TASK": "llama3.1:8b"})
    model["router"] = "MATH_TASK"
    model["replies"] += ["447", "447"]
    assert chat("What is 15% of 2400 plus 87?")[1]["meta"]["model"] == "llama3.1:8b"
    events, done = chat("What is 15% of 2400 plus 87?", think=True)
    assert model["models"][-1] == "deepseek-r1:8b" and model["think"][-1] is True and len(model["models"]) == 2


def test_ollama_receives_think_only_where_it_should(monkeypatch):
    payload = ollama_client._payload("llama3.2:3b", "p", None, 4096, True, None, None, None, True)
    assert "think" not in payload
    for think in (None, False, True):  # qwen3's slow reasoning phase is never switched on
        assert ollama_client._payload("qwen3:8b", "p", None, 4096, True, None, None, None, think)["think"] is False
    monkeypatch.setattr(config, "THINKING_MODELS", ("deepseek-r1:8b",))
    assert ollama_client._payload("deepseek-r1:8b", "p", None, 4096, True, None, None, None, True)["think"] is True


def tokens_of(events):
    """The answer text as the UI shows it: tokens after the last reset."""
    shown = []
    for e in events:
        if e["type"] == "reset":
            shown = []
        elif e["type"] == "token":
            shown.append(e["text"])
    return "".join(shown)


def step_text(events, label_part):
    """Everything streamed into the timeline step(s) whose label contains label_part."""
    ids = {e["id"] for e in events if e["type"] == "step" and label_part in e.get("label", "")}
    return "".join(e["text"] for e in events if e["type"] == "step_token" and e["id"] in ids)


def test_deep_thinking_stops_when_the_review_finds_nothing(model):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["Roses are red, violets are blue.", "NO ISSUES"]
    events, done = chat("Write a two-line poem", think=True)
    assert len(model["models"]) == 2 and "Take your time" in model["system"]
    review = model["prompts"][1]
    assert "strict reviewer" in review and "Roses are red" in review and '"Write a two-line poem"' in review
    assert tokens_of(events) == "Roses are red, violets are blue." and done["meta"]["rounds"] == 2
    assert step_text(events, "first draft") == "Roses are red, violets are blue."  # the draft is visible live
    assert step_text(events, "Reviewing version 1") == "NO ISSUES"
    assert any(e["type"] == "step" and "passed review" in e.get("label", "") for e in events)


def test_deep_thinking_revises_until_the_review_is_clean(model):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["Roses are red.", "- The poem has only one line; two were asked for.",
                         "Roses are red,\nviolets are blue.", "No issues."]
    events, done = chat("Write a two-line poem", think=True)
    assert done["reply"] == "Roses are red,\nviolets are blue." and done["meta"]["rounds"] == 4
    revise = model["prompts"][2]
    assert "Problems found in it" in revise and "only one line" in revise and "Roses are red." in revise
    assert '"Write a two-line poem"' in revise  # the original task is re-injected every round
    assert "version 2" in model["prompts"][3] and tokens_of(events) == done["reply"]  # draft 1 never reached the answer
    assert [e["phase"] for e in events if e["type"] == "phase"][-1] == "generating"


def test_deep_thinking_feeds_failed_code_checks_back_even_if_the_review_says_fine(model):
    model["router"] = "CODE_TASK"
    model["replies"] += ["```python\nprint('hi'\n```", "NO ISSUES", "```python\nprint('hi')\n```", "NO ISSUES"]
    events, done = chat("Write a python hello world program", think=True)
    assert "Automatic checks found these problems" in model["prompts"][1] and "never closed" in model["prompts"][1]
    assert "was never closed" in model["prompts"][2]  # the checker's finding is a problem to fix
    assert done["reply"] == "```python\nprint('hi')\n```"
    ids = {e["id"] for e in events if e["type"] == "step" and "Checking the code" in e.get("label", "")}
    statuses = [e["status"] for e in events if e["type"] == "step" and e["id"] in ids and e["status"] != "running"]
    assert statuses == ["warn", "done"]  # the small fix is checked, then it settles without another review
    assert len(model["models"]) == 3 and any(e.get("label") == "Version 2 settled" for e in events)


@pytest.mark.parametrize("old,new,same", [
    ("Speed = 287 / 3.75 = **76.5 km/h**", "Speed = 287 / 3.5833 = **80.1 km/h**", False),  # a real correction
    ("Steps...\nTotal = 9368\n\n**Final answer: 9368**", "Steps, clearer...\nTotal = 9368\n\n### Final answer: **9368**", True),
    ("Roses are red.", "Roses are red,\nviolets are blue.", False),
])
def test_converged(old, new, same):
    assert router_engine._converged(old, new) is same


def test_a_revision_that_breaks_the_code_is_thrown_away(model):
    model["router"] = "CODE_TASK"
    good, broken, better = "```python\nx = 1\n```", "```python\nx = (\n```", "```python\ncount = 1\n```"
    model["replies"] += [good, "- name the variable better", broken, "- still name it better", better]
    events, done = chat("Write a python variable", think=True)
    assert any(e["type"] == "step" and e.get("label") == "Kept version 1" for e in events)
    # it goes back to version 1 and tries again, told which mistake the rejected rewrite made
    assert "do not repeat it" in model["prompts"][4] and "never closed" in model["prompts"][4]
    assert done["reply"] == better and len(model["models"]) == 5  # the fixed rewrite settles at once
    model["replies"] += [good, "- name the variable better", broken, "- still", broken, "- still", broken]
    events, done = chat("Write a python variable", think=True)
    assert done["reply"] == good  # every rewrite broke it: the working version is what you get


def test_the_refine_loop_is_bounded(model):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["v1"] + [r for i in range(2, 5) for r in ("- still not great", f"v{i}")]
    events, done = chat("Write a slogan", think=True)
    assert done["reply"] == "v4" and done["meta"]["rounds"] == 1 + 2 * config.REFINE_MAX_ROUNDS
    assert not model["replies"]


def test_the_model_stays_loaded_for_a_follow_up_and_others_are_unloaded_first(model, monkeypatch):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["A poem.", "NO ISSUES"]
    chat("Write a poem", think=True)
    assert model["keep_alive"] == [config.IDLE_KEEP_ALIVE] * 2 and model["unloaded"] == []
    # the next answer needs another model: the idle one is unloaded first (one big model at a time)
    monkeypatch.setattr(ollama_client, "loaded_models", lambda: ["gemma2:9b", config.EMBED_MODEL + ":latest"])
    model["router"] = "GENERAL_TASK"
    model["replies"] += ["Hi."]
    events, done = chat("hello")
    assert model["unloaded"] == ["gemma2:9b"]  # the embedding model stays
    assert any(e.get("label") == "Unloaded the previous model" for e in events)
    monkeypatch.setattr(config, "IDLE_KEEP_ALIVE", "0")
    model["replies"] += ["Hi again."]
    chat("hello again")
    assert model["unloaded"][-1] == config.SPECIALISTS["GENERAL_TASK"]  # configured to free RAM at once


def test_fast_mode_shows_code_checks_without_another_model_call(model):
    model["router"] = "CODE_TASK"
    model["replies"] += ["```python\nprint(x)\n```"]
    events, done = chat("Write a python print statement")
    assert len(model["models"]) == 1
    assert "Undefined name `x`" in step_text(events, "Checking the code")


def test_deep_thinking_reviews_long_document_content_before_saving(model):
    content = "## Summary\n" + "Sales grew in every month of the year. " * 8
    better = content + "\n## Next steps\n- Hire two people"
    model["replies"] += [f'TOOL_CALL: create_document(filename="r.docx", title="R", content="""{content}""")',
                         f"```markdown\n{better}\n```"]
    events, done = chat("Create a Word report about sales with next steps", think=True)
    assert "You are about to save this content into r.docx" in model["prompts"][1]
    text = tools.read_docx_text(str(documents.outputs_dir() / "r.docx"))
    assert "Hire two people" in text and "```" not in text and done["reply"].startswith("Created r.docx")


def test_fast_mode_is_a_single_visible_pass(model):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["Roses are red."]
    events, done = chat("Write a two-line poem")
    assert len(model["models"]) == 1 and "Take your time" not in model["system"]
    assert "".join(e["text"] for e in events if e["type"] == "token") == "Roses are red."


def test_deep_thinking_reviews_a_tool_answer_but_not_a_created_file(model, tmp_path):
    (tmp_path / "a.txt").write_text("42")
    model["replies"] += [f"TOOL_CALL: read_file(path={str(tmp_path / 'a.txt')!r})", "The file says 42.", "NO ISSUES"]
    events, done = chat(f"read the file {tmp_path / 'a.txt'}", think=True)
    assert done["reply"] == "The file says 42." and len(model["models"]) == 3
    assert "Tool result" in model["prompts"][2]  # the review sees the tool results, so it can catch invented facts
    assert "42" in step_text(events, "Running read_file")
    model["replies"] += ['TOOL_CALL: create_document(filename="d.docx", title="T", content="hi")']
    events, done = chat("Create a Word report saying hi", think=True)
    assert done["reply"].startswith("Created d.docx")  # the file exists: nothing to review, no extra call


# --- context re-injection and follow-ups ------------------------------------------------------------------------

def test_recent_conversation_is_re_injected(model):
    model["router"] = "GENERAL_TASK"
    model["replies"] += ["Your name is Sam."]
    chat("What is my name?", history="User: My name is Sam.\nAssistant: Nice to meet you, Sam.")
    assert "Conversation so far" in model["prompts"][0] and "My name is Sam." in model["prompts"][0]
    assert model["prompts"][0].rstrip().endswith("What is my name?")


def test_short_follow_up_reuses_the_previous_route_without_the_router(model):
    model["replies"] += ["Shorter poem."]
    events, done = chat("make it shorter", prev_category="WRITING_TASK")
    assert model["router_calls"] == 0 and model["models"] == [config.SPECIALISTS["WRITING_TASK"]]
    assert done["meta"]["routed_by"] == "follow-up" and done["meta"]["router_seconds"] == 0


@pytest.mark.parametrize("message,follow_up", [
    ("make it shorter", True), ("why?", True), ("and in French?", True), ("fix that", True), ("explain more", True),
    ("Who painted the Mona Lisa?", False), ("Write a poem about the sea", False),
    ("x" * 300 + " it", False),  # long messages are new requests
])
def test_is_follow_up(message, follow_up):
    assert router_engine.is_follow_up(message) is follow_up


def test_a_tool_request_never_rides_on_a_follow_up_route(model):
    model["replies"] += ["ok"]
    events, done = chat("now make it a Word report", prev_category="WRITING_TASK")
    assert done["meta"]["category"] == "CODE_TASK" and done["meta"]["routed_by"] == "rule"


# --- loop engineering: compaction, reminders, duplicates, tool selection, warm model ----------------------------

def test_repeated_identical_tool_call_is_skipped_and_the_model_must_answer(model):
    model["replies"] += ["TOOL_CALL: get_system_info()", "TOOL_CALL: get_system_info()", "You have plenty of RAM."]
    events, done = chat("How much free RAM does this computer have right now?")
    assert sum(e["type"] == "tool" for e in events) == 1  # ran once
    assert "Skipped: you already made this exact call" in model["prompts"][2]
    assert "No more tools can be used" in model["prompts"][2] and done["reply"] == "You have plenty of RAM."


def test_older_results_are_shortened_and_the_request_is_restated(model, tmp_path):
    big = tmp_path / "big.txt"
    big.write_text("A" * 1500)
    small = tmp_path / "small.txt"
    small.write_text("tiny")
    model["replies"] += [f"TOOL_CALL: read_file(path={str(big)!r})", f"TOOL_CALL: read_file(path={str(small)!r})", "done"]
    chat("show the files and read both files in my project")
    second, third = model["prompts"][1], model["prompts"][2]
    assert "A" * 1500 in second  # newest result in full
    assert "A" * 1500 not in third and "shortened: an earlier result" in third and "tiny" in third
    for p in (second, third):
        assert 'Reminder - the user\'s request is: "show the files and read both files in my project"' in p


def test_tool_selection_lists_only_what_the_request_needs():
    docs = router_engine.select_tools("Create a Word report on testing")
    git = router_engine.select_tools("commit my changes and push")
    vague = router_engine.select_tools("help me with this")
    assert "create_document" in docs and "git_push_request" not in docs and "web_search" not in docs
    assert "git_commit" in git and "git_push_request" in git and "create_chart" not in git
    assert vague == {}  # nothing matched: a plain coding question needs no tools (saves ~40 s of prompt reading)
    assert "TOOL_CALL" not in router_engine.build_system("CODE_TASK", vague)
    assert "TOOL_CALL" in router_engine.build_system("CODE_TASK")  # no toolset given: every tool, as before
    assert len(router_engine.build_system("CODE_TASK", docs)) < len(router_engine.build_system("CODE_TASK")) * 0.8


def test_a_tool_outside_the_selection_still_works_when_called(model):
    model["replies"] += ["TOOL_CALL: get_system_info()", "ok"]
    events, done = chat("Create a Word report about my RAM usage")  # system tool group not mentioned explicitly? it is
    assert any(e["type"] == "tool" and e["name"] == "get_system_info" for e in events)


def test_keep_warm_keeps_the_model_loaded(model):
    model["replies"] += ["TOOL_CALL: get_system_info()", "ok"]
    chat("How much free RAM does this computer have right now?", keep_warm=True)
    assert set(model["keep_alive"]) == {config.WARM_KEEP_ALIVE} and model["unloaded"] == []
    model["router"] = "GENERAL_TASK"
    model["replies"] += ["hi"]
    chat("hello", keep_warm=True)
    assert model["keep_alive"][-1] == config.WARM_KEEP_ALIVE


# --- routing -----------------------------------------------------------------------------------------------

@pytest.mark.parametrize("message", [
    "Create a Word report on unit testing", "Make an Excel spreadsheet of my expenses", "Generate a PDF summary of cats",
    "Build a PowerPoint presentation about Python", "make a bar chart of sales", "Export this as csv",
    "commit my changes", "push the branch to github", "stage the files in my project", "git commit with message x",
])
def test_document_and_git_requests_go_to_the_tool_model(message):
    assert router_engine.rule_category(message) == "CODE_TASK"


@pytest.mark.parametrize("message", [
    "How does a PDF store fonts?", "Write a poem about the sea", "push yourself harder",
    "Tell me a joke about spreadsheets",
])
def test_ordinary_questions_are_left_to_the_router(message):
    assert router_engine.rule_category(message) is None


# --- git and the push dialog, in chat ---------------------------------------------------------------------------

def git(cwd, *args):
    return subprocess.run(["git", *args], cwd=cwd, capture_output=True, text=True, check=True).stdout


@pytest.fixture
def repo(tmp_path, monkeypatch):
    for k, v in {"GIT_AUTHOR_NAME": "T", "GIT_AUTHOR_EMAIL": "t@example.org",
                 "GIT_COMMITTER_NAME": "T", "GIT_COMMITTER_EMAIL": "t@example.org"}.items():
        monkeypatch.setenv(k, v)
    bare, work = tmp_path / "remote.git", tmp_path / "work"
    work.mkdir()
    git(tmp_path, "init", "--bare", "-b", "main", str(bare))
    git(work, "init", "-b", "feature")
    git(work, "remote", "add", "origin", str(bare))
    (work / "app.py").write_text("print('hi')\n")
    git(work, "add", "app.py")
    git(work, "commit", "-m", "first")
    con.save_settings({"allowed_remote_hosts": ["github.com", "local"]})
    return work, bare


def on_remote(bare, branch="feature"):
    return subprocess.run(["git", "--git-dir", str(bare), "rev-parse", "--verify", f"refs/heads/{branch}"],
                          capture_output=True).returncode == 0


def test_stage_commit_and_push_from_chat_after_the_users_click(model, repo, monkeypatch):
    work, bare = repo
    (work / "notes.py").write_text("NOTES = 1\n")
    model["replies"] += [
        f'TOOL_CALL: git_add(repo_path={str(work)!r}, paths="notes.py")',
        f'TOOL_CALL: git_commit(repo_path={str(work)!r}, message="add notes")',
        f'TOOL_CALL: git_push_request(repo_path={str(work)!r})',
    ]
    monkeypatch.setattr(approvals, "wait", lambda aid, cancel=None, timeout=0: True)
    events, done = chat(f"commit my changes in {work} and push the branch")
    ask = next(e for e in events if e["type"] == "push_approval")
    assert ask["branch"] == "feature" and ask["violations"] == [] and "add notes" in ask["review"]
    assert done["reply"].startswith("Pushed.") and on_remote(bare) and done["meta"]["rounds"] == 3
    assert "add notes" in git(str(bare), "log", "--oneline", "feature")


def test_a_garbled_repo_path_falls_back_to_the_repo_used_earlier(model, repo, monkeypatch):
    """Real failure: on its third git call the model truncated the long repository path."""
    work, bare = repo
    (work / "n.py").write_text("n = 1\n")
    broken = str(work)[:-6]  # what the model actually produced: a cut-off path
    model["replies"] += [
        f'TOOL_CALL: git_add(repo_path={str(work)!r}, paths="n.py")',
        f'TOOL_CALL: git_commit(repo_path={str(work)!r}, message="add n")',
        f'TOOL_CALL: git_push_request(repo_path={broken!r})',
    ]
    monkeypatch.setattr(approvals, "wait", lambda aid, cancel=None, timeout=0: True)
    events, done = chat(f"commit my changes in {work} and push the branch")
    assert done["reply"].startswith("Pushed.") and on_remote(bare)


def test_a_wrong_path_without_an_earlier_repo_is_not_guessed(model, repo, monkeypatch):
    work, bare = repo
    model["replies"] += [f'TOOL_CALL: git_push_request(repo_path={str(work)[:-6]!r})', "That folder does not exist."]
    events, done = chat("push the branch to github")
    assert not any(e["type"] == "push_approval" for e in events) and not on_remote(bare)


def test_declined_push_pushes_nothing(model, repo, monkeypatch):
    work, bare = repo
    model["replies"] += [f'TOOL_CALL: git_push_request(repo_path={str(work)!r})']
    monkeypatch.setattr(approvals, "wait", lambda aid, cancel=None, timeout=0: False)
    events, done = chat("push the branch to github")
    assert "declined" in done["reply"] and not on_remote(bare)
    assert "push_declined" in [e["event"] for e in con.read_audit()]


def test_a_click_cannot_push_past_the_constitution_without_an_override(model, repo, monkeypatch):
    work, bare = repo
    (work / "s.py").write_text(f"K = '{AWS}'\n")
    git(work, "add", "s.py")
    git(work, "commit", "-m", "oops")
    model["replies"] += [f'TOOL_CALL: git_push_request(repo_path={str(work)!r})', "It was blocked."]
    monkeypatch.setattr(approvals, "wait", lambda aid, cancel=None, timeout=0: True)
    events, done = chat("push the branch to github")
    tool = next(e for e in events if e["type"] == "tool")
    assert tool["result"].startswith("Blocked by the constitution") and not on_remote(bare)


def test_an_override_from_the_dialog_travels_with_the_click(model, repo):
    work, bare = repo
    (work / "s.py").write_text(f"K = '{AWS}'\n")
    git(work, "add", "s.py")
    git(work, "commit", "-m", "oops")
    model["replies"] += [f'TOOL_CALL: git_push_request(repo_path={str(work)!r})']

    def user_clicks():
        for _ in range(200):
            if approvals.pending():
                approvals.decide(approvals.pending()[0], True, ["secrets_in_push", "not_a_rule"])
                return
            time.sleep(0.05)
    threading.Thread(target=user_clicks, daemon=True).start()
    events, done = chat("push the branch to github")
    assert done["reply"].startswith("Pushed.") and on_remote(bare)
    assert [e for e in con.read_audit() if e["event"] == "push_done"][0]["overrides_used"] == ["secrets_in_push"]


# --- lifting the whole constitution -----------------------------------------------------------------------------------

def test_lift_everything_switches_every_rule_off_for_a_while(monkeypatch, tmp_path):
    (tmp_path / ".env").write_text("TOKEN=x")
    assert "refusing" in tools.read_file(str(tmp_path / ".env"))
    assert documents.create_document("k.docx", "K", f"key {AWS}").startswith("Error: Refused")
    assert "Refused by the constitution" in tools.web_search(f"key {AWS}")
    con.overrides.grant_all(5)
    assert con.overrides.lifted_seconds() > 0 and all(not con.enforced(r["id"]) for r in con.RULES)
    assert tools.read_file(str(tmp_path / ".env")) == "TOKEN=x"
    assert documents.create_document("k.docx", "K", f"key {AWS}").startswith("Created")
    assert tools._refuse_secret("search query", f"key {AWS}") is None  # the guard steps aside (no network in tests)
    real = time.time
    monkeypatch.setattr(time, "time", lambda: real() + 6 * 60)
    assert con.overrides.lifted_seconds() == 0 and con.enforced("secrets_in_documents")  # it ends by itself
    events = [e["event"] for e in con.read_audit()]
    assert "constitution_lifted" in events


def test_lift_everything_lets_a_secret_push_through_but_still_needs_the_click(model, repo, monkeypatch):
    work, bare = repo
    (work / "s.py").write_text(f"K = '{AWS}'\n")
    git(work, "add", "s.py")
    git(work, "commit", "-m", "oops")
    con.overrides.grant_all(10)
    review = tools.git_push_review(str(work))
    assert [v["rule"] for v in review["violations"]] == ["secrets_in_push"] and review["violations"][0]["overridden"]
    model["replies"] += [f'TOOL_CALL: git_push_request(repo_path={str(work)!r})']
    asked = []
    monkeypatch.setattr(approvals, "wait", lambda aid, cancel=None, timeout=0: asked.append(aid) or True)
    events, done = chat("push the branch to github")
    assert asked and done["reply"].startswith("Pushed.") and on_remote(bare)


def test_restore_ends_the_lift_early():
    con.overrides.grant_all(30)
    con.overrides.revoke()
    assert con.overrides.lifted_seconds() == 0 and con.enforced("secrets_in_push")
    assert [e["event"] for e in con.read_audit()][:1] == ["constitution_restored"]
    for bad in (0, 61, "10", True):
        with pytest.raises(ValueError):
            con.overrides.grant_all(bad)


def test_api_lift_and_restore(client):
    assert client.post("/api/constitution/lift-all", json={"minutes": 120}).status_code == 400
    r = client.post("/api/constitution/lift-all", json={"minutes": 15})
    assert r.status_code == 200 and 800 < r.json()["lifted_seconds"] <= 900
    assert client.get("/api/status").json()["constitution_lifted_seconds"] > 0
    assert client.get("/api/constitution").json()["lifted_seconds"] > 0
    assert client.post("/api/constitution/lift-all", json={"minutes": 5}, headers={"Origin": "https://evil.example"}).status_code == 403
    client.delete("/api/constitution/overrides")
    assert client.get("/api/constitution").json()["lifted_seconds"] == 0


# --- approvals broker -----------------------------------------------------------------------------------------------

def test_approval_broker():
    aid = approvals.create()
    assert approvals.decide(aid, True, ["protected_branches", "bogus"]) is True
    assert approvals.wait(aid, timeout=2) is True and approvals.take_overrides(aid) == ["protected_branches"]
    assert approvals.decide(aid, True) is False  # consumed
    assert approvals.wait(approvals.create(), timeout=0.6) is False  # nobody answered
    cancel = threading.Event()
    cancel.set()
    assert approvals.wait(approvals.create(), cancel, timeout=5) is False  # tab closed
    assert approvals_mod.request_push.__doc__  # the only push path the chat has


# --- API ------------------------------------------------------------------------------------------------------------------

@pytest.fixture
def client():
    from fastapi.testclient import TestClient
    import main
    with TestClient(main.app, base_url="http://127.0.0.1:8420") as c:
        yield c


def sse(response):
    return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]


def test_api_chat_creates_a_document_and_serves_it(client, model):
    model["replies"] += ['TOOL_CALL: create_document(filename="plan.docx", title="Plan", content="## Steps\\n- one")']
    events = sse(client.post("/api/chat/stream", json={"message": "Create a Word report with my plan"}))
    done = events[-1]
    assert done["type"] == "done" and done["meta"]["files"] == ["plan.docx"]
    assert any(e["type"] == "files" for e in events)
    got = client.get("/api/outputs/download", params={"path": "plan.docx"})
    assert got.status_code == 200 and got.content[:2] == b"PK"
    assert client.get("/api/outputs").json()["files"][0]["name"] == "plan.docx"
    for bad, code in (("../main.py", 400), ("..\\main.py", 400), ("C:\\Windows\\win.ini", 400), ("nope.docx", 404)):
        assert client.get("/api/outputs/download", params={"path": bad}).status_code == code
    msgs = client.get(f"/api/chats/{done['chat']['id']}/messages").json()["messages"]
    assert msgs[-1]["meta"]["files"] == ["plan.docx"]  # download buttons survive a reload


def test_api_push_dialog_round_trip(client, model, repo):
    work, bare = repo
    model["replies"] += [f'TOOL_CALL: git_push_request(repo_path={str(work)!r})']
    assert client.post("/api/approve", json={"id": "nope", "approved": True}).status_code == 404

    def user_clicks():
        for _ in range(200):
            if approvals.pending():
                aid = approvals.pending()[0]
                r = client.post("/api/approve", json={"id": aid, "approved": True, "overrides": ["secrets_in_push"]})
                assert r.status_code == 400  # a core override without the typed word is refused...
                client.post("/api/approve", json={"id": aid, "approved": True})  # ...a plain click is fine here
                return
            time.sleep(0.05)
    threading.Thread(target=user_clicks, daemon=True).start()
    events = sse(client.post("/api/chat/stream", json={"message": "push the branch to github"}))
    assert any(e["type"] == "push_approval" for e in events) and events[-1]["reply"].startswith("Pushed.")
    assert on_remote(bare)


def test_api_blocking_chat_never_pushes(client, model, repo):
    work, bare = repo
    model["replies"] += [f'TOOL_CALL: git_push_request(repo_path={str(work)!r})']
    r = client.post("/api/chat", json={"message": "push the branch to github"}).json()
    assert "declined" in r["reply"] and not on_remote(bare)


def test_api_follow_up_reuses_route_and_sees_real_history(client, model):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["Roses are red, violets are blue, my name is Sam.", "Roses are red."]
    first = sse(client.post("/api/chat/stream", json={"message": "Write a two-line poem that mentions Sam"}))[-1]
    second = sse(client.post("/api/chat/stream", json={"message": "make it shorter", "chat_id": first["chat"]["id"]}))[-1]
    assert model["router_calls"] == 1 and second["meta"]["routed_by"] == "follow-up"
    assert "Write a two-line poem that mentions Sam" in model["prompts"][1] and "violets are blue" in model["prompts"][1]


def test_api_passes_thinking_and_keep_warm(client, model):
    model["router"] = "MATH_TASK"
    model["replies"] += ["4", "4", "NO ISSUES"]
    fast = sse(client.post("/api/chat/stream", json={"message": "2+2?"}))[-1]
    deep = sse(client.post("/api/chat/stream", json={"message": "2+2?", "think": True, "keep_warm": True}))[-1]
    assert model["think"] == [False, True, True] and fast["meta"]["thinking"] == "fast"
    assert deep["meta"]["thinking"] == "deep" and deep["meta"]["warm"] is True and deep["meta"]["rounds"] == 2


def test_api_saves_the_timeline_with_the_answer(client, model):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["A poem.", "NO ISSUES"]
    events = sse(client.post("/api/chat/stream", json={"message": "Write a poem", "think": True}))
    assert {e["type"] for e in events} >= {"step", "step_token", "token", "done"}
    msgs = client.get(f"/api/chats/{events[-1]['chat']['id']}/messages").json()["messages"]
    labels = [s["label"] for s in msgs[-1]["meta"]["steps"]]
    assert labels[0] == "Choosing a specialist" and "Writing a first draft" in labels and "Reviewing version 1" in labels
    assert all(s["status"] != "running" and "seconds" in s for s in msgs[-1]["meta"]["steps"])


def test_downloads_never_render_inside_the_app(client):
    documents.create_document("page.html", "P", "<!DOCTYPE html><html><body><script>alert(1)</script></body></html>")
    got = client.get("/api/outputs/download", params={"path": "page.html"})
    assert got.headers["content-disposition"].startswith("attachment")
    assert got.headers["content-security-policy"] == "sandbox" and got.headers["x-content-type-options"] == "nosniff"


# --- design / HTML -----------------------------------------------------------------------------------------------

def test_html_documents_keep_a_page_or_build_one_from_markdown():
    page = "```html\n<!DOCTYPE html><html><head><title>X</title></head><body><h1>Hi</h1></body></html>\n```"
    assert documents.create_document("site.html", "Site", page).startswith("Created site.html")
    saved = (documents.outputs_dir() / "site.html").read_text(encoding="utf-8")
    assert saved.startswith("<!DOCTYPE html>") and "```" not in saved
    documents.create_document("notes.html", "Notes <b>", "## Plan\n- one & two\n\n| a | b |\n| 1 | 2 |")
    built = (documents.outputs_dir() / "notes.html").read_text(encoding="utf-8")
    assert "<h2>Plan</h2>" in built and "<li>one &amp; two</li>" in built and "<td>1</td>" in built
    assert "Notes &lt;b&gt;" in built and 'name="viewport"' in built
    assert documents.requested_outputs("Save the landing page as an HTML file") == [".html"]


def test_design_requests_get_the_design_prompt(model):
    model["router"] = "CODE_TASK"
    model["replies"] += ["```html\n<!DOCTYPE html><html><head><title>T</title></head><body></body></html>\n```"]
    events, done = chat("Design a landing page for my bakery")
    assert "ONE complete, self-contained HTML file" in model["system"] and done["meta"]["design"] is True
    model["replies"] += ["def f(): pass"]
    chat("Write a python function that adds two numbers")
    assert "self-contained HTML" not in model["system"]
    assert router_engine.build_system("WRITING_TASK", design=True) == router_engine.SYSTEM_PROMPTS["WRITING_TASK"]


@pytest.mark.parametrize("text,issues", [
    ("**Average speed = 287 / 3.5833 ≈ 79.9 km/h**", ["287 / 3.5833 is 80.0938, not 79.9"]),  # the real slip
    ("Average speed = 287 km / 3.5833 hours ≈ 80.1 km/h", []),
    ("Total = 6633 + 4100 - 1365 = 10733 - 1365 = 9368", []),  # a chain of equal steps
    ("Total = 6633 + 4100 - 1365 = 9468", ["6633 + 4100 - 1365 is 9368, not 9468"]),
    ("35 / 60 = 0.5833 hours, total 3 + 0.5833 = 3.5833", []),
    ("Leaves 9:40, arrives 13:15 - 9:40 later", []),  # times are not arithmetic
    ("```python\nx = 2 + 2 = 5\n```", []),  # code blocks are left to the code checks
    ("Probability = 5 / 36 ≈ 0.139", []),
    ("Average speed = (287)/(3.5833) = 79.8", ["(287)/(3.5833) is 80.0938, not 79.8"]),  # real run 2
    (r"$$ \text{Speed} = \frac{287}{3.5833} = 79.8 $$", ["(287)/(3.5833) is 80.0938, not 79.8"]),  # LaTeX
    ("(2 + 3) * 4 = 20", []),
])
def test_arithmetic_is_recomputed(text, issues):
    import checks
    found = checks.check_arithmetic(text)[1]
    assert [i.split("Arithmetic: ")[1] for i in found] == issues


@pytest.mark.parametrize("latex,plain", [
    ("$$\n\\text{Total sum} = 6633 + 4100 - 1365 = 9368\n$$", "Total sum = 6633 + 4100 - 1365 = 9368"),
    ("$$\n\\boxed{\\frac{5}{36}}\n$$", "**(5)/(36)**"),
    ("Speed $\\approx 80.1$ km/h and $3 \\times 4 = 12$", "Speed ≈ 80.1 km/h and 3 × 4 = 12"),
    ("It costs $5 and $10 total", "It costs $5 and $10 total"),  # money is not math
    ("```python\nprint('$x$ \\\\frac')\n```", "```python\nprint('$x$ \\\\frac')\n```"),  # code untouched
])
def test_plain_math(latex, plain):
    assert router_engine.plain_math(latex).strip() == plain


def test_math_replies_reach_the_user_as_plain_text(model):
    model["router"] = "MATH_TASK"
    model["replies"] += ["Answer:\n$$\n\\boxed{9368}\n$$"]
    reply = chat("Sum of multiples of 3 or 5 up to 200?")[1]["reply"]
    assert "**9368**" in reply and "$" not in reply and "boxed" not in reply


def test_a_wrong_calculation_blocks_the_early_stop_and_is_fed_back(model):
    model["router"] = "MATH_TASK"
    model["replies"] += ["Speed = 287 / 3.5833 ≈ 79.9 km/h", "NO ISSUES", "Speed = 287 / 3.5833 ≈ 80.1 km/h", "NO ISSUES"]
    events, done = chat("Average speed of a train going 287 km in 3 h 35 min?", think=True)
    assert "287 / 3.5833 is 80.0938, not 79.9" in model["prompts"][1] and done["reply"].endswith("80.1 km/h")


def test_checks_understand_a_language_name_on_its_own_line():
    """qwen2.5-coder wrote ```, then 'python' on the next line (benchmark)."""
    import checks
    assert checks.code_blocks("```\npython\nx = 1\n```") == [("python", "x = 1\n")]
    assert checks.run_checks("```\npython\nprint(1)\n```") == {"checked": ["python"], "issues": []}


@pytest.mark.parametrize("text,looping", [
    ("<img src='data:image/svg+xml;base64,PHN2Zy" + "MTkgMTBM" * 40, True),  # the real loop that ran for 8 minutes
    ("Intro. " + "The same sentence again and again, forever and ever. " * 6, True),
    ("x = [0] * 100\n" + "- item\n" * 5 + "=" * 60, False),
    ("def f():\n" + "    " * 80 + "return 1", False),  # indentation is not a loop
    ("Short answer.", False),
])
def test_repeating_tail(text, looping):
    assert (router_engine.repeating_tail(text) > 0) is looping


def test_a_looping_generation_is_cut_off_and_the_useful_part_kept(model):
    model["router"] = "CODE_TASK"
    model["replies"] += ["Here is the page:\n<img src='data:" + "MTkgMTBM" * 400]
    events, done = chat("Write a python print statement")
    assert done["reply"].startswith("Here is the page:") and "repeating the same text" in done["reply"]
    assert len(done["reply"]) < 400 and tokens_of(events) == done["reply"]
    assert any(e["type"] == "step" and e["status"] == "warn" and "repetition" in e.get("detail", "") for e in events)


# --- prewarm while typing ---------------------------------------------------------------------------------------

@pytest.mark.parametrize("message,prev,deep,think,expected", [
    ("Create a Word report on testing", None, False, False, (config.SPECIALISTS["CODE_TASK"], config.TOOL_NUM_CTX)),
    ("explain recursion", "CODE_TASK", False, False, (config.SPECIALISTS["CODE_TASK"], config.DEFAULT_NUM_CTX)),
    ("make it shorter", "WRITING_TASK", False, False, (config.SPECIALISTS["WRITING_TASK"], config.DEFAULT_NUM_CTX)),
    ("why?", "MATH_TASK", False, False, ("qwen3:8b", config.DEFAULT_NUM_CTX)),
    ("why?", "SECURITY_TASK", False, True, (config.SPECIALISTS["SECURITY_TASK"], config.DEFAULT_NUM_CTX)),
    ("Who painted the Mona Lisa?", None, False, False, (config.ROUTER_MODEL, config.ROUTER_NUM_CTX)),
    ("anything at all", None, True, False, None),
])
def test_prewarm_target_matches_the_route(message, prev, deep, think, expected):
    target = router_engine.prewarm_target(message, prev, deep, think)
    assert (target[:2] if target else None) == expected


def test_api_prewarm_loads_only_when_it_is_safe(client, monkeypatch):
    import main
    loaded, started = [], threading.Event()
    monkeypatch.setattr(main.psutil, "virtual_memory", lambda: type("VM", (), {"available": 12 * 2**30})())
    monkeypatch.setattr(ollama_client, "loaded_models", lambda: list(loaded))
    monkeypatch.setattr(ollama_client, "preload", lambda m, ctx, ka: started.set() or loaded.append(m) or True)
    r = client.post("/api/prewarm", json={"message": "Who painted the Mona Lisa?"}).json()
    assert r == {"started": True, "model": config.ROUTER_MODEL} and started.wait(5)
    assert client.post("/api/prewarm", json={"message": "Who painted it?"}).json()["reason"] == "already loaded"
    assert client.post("/api/prewarm", json={"message": "x", "deep": True}).json()["started"] is False
    main.GEN_LOCK.acquire()
    try:
        assert client.post("/api/prewarm", json={"message": "Create a Word report"}).json()["reason"] == "busy"
    finally:
        main.GEN_LOCK.release()
    monkeypatch.setattr(main.psutil, "virtual_memory", lambda: type("VM", (), {"available": 2**30})())
    assert "GB free" in client.post("/api/prewarm", json={"message": "Create a Word report"}).json()["reason"]


def test_cpu_temperature_prefers_the_cpu_zone():
    import sysmon
    assert sysmon.cpu_temperature({"_tz.batz": 33.0, "_tz.cpuz": 72.9, "_tz.chgz": 51.0}) == 72.9
    assert sysmon.cpu_temperature({"_tz.a": 40.0, "_tz.b": 55.5}) == 55.5  # no CPU zone: the hottest
    assert sysmon.cpu_temperature({}) is None


def test_api_system_live_meters(client):
    s = client.get("/api/system/live").json()
    assert set(s) == {"cpu_percent", "ram_used_gb", "ram_total_gb", "ram_percent", "model_ram_gb", "cpu_temp_c"}
    assert 0 <= s["cpu_percent"] <= 100 and 0 < s["ram_used_gb"] <= s["ram_total_gb"]
    assert s["cpu_temp_c"] is None or 0 < s["cpu_temp_c"] < 120


def test_agent_endpoints_are_gone(client):
    for path in ("/api/agent/run", "/api/agent/stream", "/api/agent/approve"):
        assert client.post(path, json={}).status_code in (404, 405)
    assert not (ROOT / "agent_engine.py").exists()
    assert "Agent" not in client.get("/").text.split("<main>")[1].split("</div>")[0]


def test_a_follow_up_keeps_the_tool_list_order_so_the_prompt_cache_stays_valid(model):
    model["router"] = "CODE_TASK"
    model["replies"] += ["ok"]
    first = chat("Show me the files in my project folder")[1]["meta"]["toolset"]
    model["replies"] += ["ok"]
    events, done = chat("Now show me the git status of my project", prev_category="CODE_TASK", prev_tools=first)
    second = done["meta"]["toolset"]
    assert second[:len(first)] == first and "git_status" in second[len(first):]
    s1 = router_engine.build_system("CODE_TASK", {n: tools.MODEL_TOOLS[n] for n in first})
    assert model["system"].startswith(s1)  # the whole first system prompt is a prefix of the second


def test_a_file_nobody_asked_for_is_not_created(model):
    """Real run: asked only for the git status, the model also made git_status_report.docx."""
    model["replies"] += ['TOOL_CALL: create_document(filename="status.docx", title="S", content="clean")',
                         "Your working tree is clean."]
    events, done = chat("show me the git status of my project")
    assert not (documents.outputs_dir() / "status.docx").exists() and done["reply"] == "Your working tree is clean."
    assert "did not ask for a file" in model["prompts"][1]
    assert any(e.get("label") == "Skipped create_document" for e in events)
    model["replies"] += ['TOOL_CALL: create_document(filename="status.docx", title="S", content="clean")']
    events, done = chat("show me the git status and save it as a Word report")
    assert done["reply"].startswith("Created status.docx")  # asked for: made


@pytest.fixture
def chat_api(model, monkeypatch):
    """Fake /api/chat next to the fake /api/generate of `model`; records the messages it was sent."""
    sent = []

    def stream_chat(model_name, messages, num_ctx=0, timeout=0, temperature=None, cancel=None, keep_alive=None,
                    think=None):
        sent.append(messages)
        model["models"].append(model_name)
        text = model["replies"].pop(0)
        yield ("response", text)
    monkeypatch.setattr(ollama_client, "stream_chat", stream_chat)
    monkeypatch.setattr(router_engine, "_CONVERSATIONS", {})
    return sent


def test_a_follow_up_continues_the_cached_conversation(model, chat_api, monkeypatch):
    model["router"] = "WRITING_TASK"
    writer = config.SPECIALISTS["WRITING_TASK"]
    model["replies"] += ["Roses are red, violets are blue."]
    first = chat("Write a two-line poem", conversation_key="c1")[1]
    monkeypatch.setattr(ollama_client, "loaded_models", lambda: [writer])
    model["replies"] += ["Roses are red."]
    events, done = chat("make it shorter", conversation_key="c1", prev_category="WRITING_TASK")
    msgs = chat_api[-1]
    assert [m["role"] for m in msgs] == ["system", "user", "assistant", "user"]
    assert msgs[1]["content"] == model["prompts"][0] and msgs[2]["content"] == first["reply"]
    assert msgs[3]["content"] == "User request:\nmake it shorter" and done["reply"] == "Roses are red."
    assert any(e.get("label") == "Continuing the conversation" for e in events)
    model["replies"] += ["Red."]
    chat("even shorter", conversation_key="c1", prev_category="WRITING_TASK")
    assert len(chat_api[-1]) == 6  # the conversation keeps growing while the model holds it


def test_no_continuation_when_it_would_not_be_cached_or_could_change_behaviour(model, chat_api, monkeypatch):
    model["router"] = "WRITING_TASK"
    model["replies"] += ["A poem."]
    chat("Write a poem", conversation_key="c2")
    model["replies"] += ["Shorter."]
    chat("make it shorter", conversation_key="c2", prev_category="WRITING_TASK")  # model not loaded any more
    assert chat_api == []
    monkeypatch.setattr(ollama_client, "loaded_models", lambda: [config.SPECIALISTS["WRITING_TASK"]])
    model["replies"] += ["Deep.", "NO ISSUES"]
    chat("make it deeper", conversation_key="c2", prev_category="WRITING_TASK", think=True)
    assert chat_api == [] and "c2" not in router_engine._CONVERSATIONS  # deep thinking: fresh prompts, state dropped


def test_a_revision_that_drops_the_working_is_rejected(model):
    """Real run: each revision was only '**Final Answer:** 79.8 km/h', so the arithmetic check had nothing to check."""
    model["router"] = "MATH_TASK"
    work = ("We need the average speed: the distance divided by the time taken for the whole trip, in hours.\n"
            "Step 1: 13:15 - 9:40 is 3 h 35 min.\nStep 2: 35 / 60 = 0.5833 h, so 3 + 0.5833 = 3.5833 h.\n"
            "Step 3: speed = 287 / 3.5833 = 79.9 km/h.\nStep 4: check the units.\nStep 5: round.\nStep 6: done.\n"
            "Final answer: 79.9 km/h")
    fixed = work.replace("79.9", "80.1")
    model["replies"] += [work, "NO ISSUES", "**Final Answer:** 79.8 km/h", "**Final Answer:** 79.8 km/h"]
    events, done = chat("A train covers 287 km from 9:40 to 13:15. Average speed?", think=True)
    assert done["reply"].endswith("79.9 km/h") and "Step 1" in done["reply"]  # the complete version is kept
    assert sum(e.get("label", "").startswith("Rejected version") for e in events) == 2
    assert router_engine.DROPPED_WORK in model["prompts"][3]
    model["replies"] += [work, "NO ISSUES", "**Final Answer:** 79.8 km/h", fixed, "NO ISSUES"]
    events, done = chat("A train covers 287 km from 9:40 to 13:15. Average speed?", think=True)
    assert done["reply"] == fixed  # asked again, it wrote the full corrected answer


@pytest.mark.parametrize("text,issue", [
    ("From 9:40 to 13:15 is 4 hours 15 minutes.", "is 3 h 35 min (3.583 hours), not 4 h 15 min"),
    ("9:40 AM to 1:15 PM = 4.25 hours", "is 3 h 35 min (3.583 hours), not 4 h 15 min"),
    ("From 9:40 to 13:15 is 3 hours and 35 minutes.", None),
    ("13:15 - 9:40 = 3 h 35 min", None),
    ("between 11:50 pm and 12:20 am: 30 minutes", None),  # past midnight
    ("the trip from 9:40 to 13:15 takes about 3.6 hours", None),  # rounding is fine
    ("Meeting 10:00 to 11:00. Then lunch.", None),  # no duration stated
])
def test_time_spans_are_recomputed(text, issue):
    import checks
    found = checks.check_time_spans(text)[1]
    if issue:
        assert len(found) == 1 and issue in found[0]
    else:
        assert found == []


def test_an_arithmetic_slip_in_a_rewrite_does_not_bring_back_a_wrong_draft(model):
    """Real runs: the draft had a wrong time span (no check sees it), the review caught it, the rewrite fixed it
    but slipped on a division; going back to the draft kept the wrong answer. Now the slip is fixed next round."""
    model["router"] = "MATH_TASK"
    wrong = "Time is 3.75 h.\n" + "Speed = 287 / 3.75 = 76.53 km/h.\n" * 2 + "Final answer: 76.5 km/h"
    slip = "Time is 3.5833 h.\n" + "Speed = 287 / 3.5833 = 79.8 km/h.\n" * 2 + "Final answer: 79.8 km/h"
    right = slip.replace("79.8", "80.1")
    model["replies"] += [wrong, "- 9:40 to 13:15 is 3 h 35 min, not 3.75 h", slip, "NO ISSUES", right, "NO ISSUES"]
    events, done = chat("A train covers 287 km from 9:40 to 13:15. Average speed?", think=True)
    assert not any(e.get("label", "").startswith("Kept version") for e in events)
    assert "is 80.0938, not 79.8" in model["prompts"][4] and done["reply"] == right
