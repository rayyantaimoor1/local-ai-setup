"""Offline test suite: a fake Ollama stands in for the real models."""
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
# (conftest.py has already pointed the database, constitution settings and audit log at a temp folder)

import config  # noqa: E402
import ollama_client  # noqa: E402
import router_engine  # noqa: E402
import tools  # noqa: E402
from graph_memory import MemoryManager  # noqa: E402
from ollama_client import OllamaError, ThinkFilter  # noqa: E402


# --- fake Ollama ---------------------------------------------------------------

class FakeOllama:
    def __init__(self):
        self.handler = lambda model, prompt, system: "ok"
        self.pulled = [config.ROUTER_MODEL, *config.SPECIALISTS.values(), config.ESCALATION_MODEL]
        self.calls = []
        self.unloaded = []
        self.keep_alives = []

    def text(self, model, prompt, system):
        self.calls.append(model)
        return self.handler(model, prompt, system)


@pytest.fixture
def fake(monkeypatch):
    f = FakeOllama()

    def stream(model, prompt, system=None, num_ctx=0, timeout=0, temperature=None, cancel=None, keep_alive=None,
               think=None):
        f.keep_alives.append(keep_alive)
        text = f.text(model, prompt, system)
        for i in range(0, len(text), 7):
            if cancel is not None and cancel.is_set():
                return
            yield ("response", text[i:i + 7])

    def generate(model, prompt, system=None, num_ctx=0, timeout=0, temperature=None, num_predict=None):
        return ollama_client.strip_think(f.text(model, prompt, system))

    monkeypatch.setattr(ollama_client, "stream_generate", stream)
    monkeypatch.setattr(ollama_client, "generate", generate)
    monkeypatch.setattr(ollama_client, "list_models", lambda: list(f.pulled))
    monkeypatch.setattr(ollama_client, "unload", lambda model: f.unloaded.append(model))
    monkeypatch.setattr(router_engine, "_trim", lambda: {"triggered": False, "freed_mb": 0})
    return f


def run_chat(message, context="", deep=False):
    return list(router_engine.process_request(message, context, deep))


# --- ollama_client ----------------------------------------------------------------

def test_think_filter_across_chunks():
    f = ThinkFilter()
    out = "".join(f.feed(c) for c in ["Hel", "lo <thi", "nk>secret</th", "ink> world", "<", "x"]) + f.flush()
    assert out == "Hello  world<x"


def test_strip_think():
    assert ollama_client.strip_think("<think>a</think>Answer") == "Answer"
    assert ollama_client.strip_think("Answer<think>never closed") == "Answer"


# --- router -----------------------------------------------------------------------------

def test_parse_category():
    assert router_engine.parse_category("CODE_TASK") == "CODE_TASK"
    assert router_engine.parse_category("I think MATH_TASK, not CODE_TASK") == "MATH_TASK"
    assert router_engine.parse_category("no idea") == "GENERAL_TASK"
    assert router_engine.parse_category("") == "GENERAL_TASK"


def test_routes_to_specialist_and_streams(fake):
    fake.handler = lambda m, p, s: "WRITING_TASK" if m == config.ROUTER_MODEL else "The answer is 42."
    events = run_chat("write me a haiku")
    kinds = [e["type"] for e in events]
    assert kinds[0] == "phase" and events[0]["phase"] == "classifying"
    assert "".join(e["text"] for e in events if e["type"] == "token") == "The answer is 42."
    done = events[-1]
    assert done["type"] == "done" and done["meta"]["category"] == "WRITING_TASK"
    assert done["meta"]["model"] == config.SPECIALISTS["WRITING_TASK"]
    assert fake.calls == [config.ROUTER_MODEL, config.SPECIALISTS["WRITING_TASK"]]


def test_math_uses_qwen3_in_both_modes(fake):
    fake.handler = lambda m, p, s: "MATH_TASK" if m == config.ROUTER_MODEL else ("NO ISSUES" if "strict reviewer" in p else "42")
    assert run_chat("what is 6 times 7")[-1]["meta"]["model"] == "qwen3:8b"
    deep = list(router_engine.process_request("what is 6 times 7", "", False, think=True))[-1]
    assert deep["meta"]["model"] == "qwen3:8b" and deep["reply"] == "42"


def test_garbage_router_output_falls_back_to_general(fake):
    fake.handler = lambda m, p, s: "banana" if m == config.ROUTER_MODEL else "hi"
    assert run_chat("hello")[-1]["meta"]["category"] == "GENERAL_TASK"


def test_tool_call_is_held_back_and_executed(fake):
    def handler(model, prompt, system):
        if model == config.ROUTER_MODEL:
            return "CODE_TASK"
        if "Tool result" in prompt:
            return "Your project has config.py and main.py."
        return 'TOOL_CALL: get_directory_tree(path=".")'
    fake.handler = handler
    events = run_chat("show me the files in my project")
    streamed = "".join(e["text"] for e in events if e["type"] == "token")
    assert "TOOL_CALL" not in streamed
    assert streamed == "Your project has config.py and main.py."
    assert any(e.get("phase") == "running_tool" for e in events)
    assert events[-1]["meta"]["tool"] == "get_directory_tree"
    assert fake.calls.count(config.SPECIALISTS["CODE_TASK"]) == 2  # two generations per tool request


def test_deep_mode_skips_router(fake):
    fake.handler = lambda m, p, s: "careful answer"
    done = run_chat("hard question", deep=True)[-1]
    assert done["meta"]["model"] == config.ESCALATION_MODEL and done["meta"]["category"] == "DEEP"
    assert fake.calls == [config.ESCALATION_MODEL]


def test_ollama_error_becomes_error_event(fake, monkeypatch):
    def boom(*a, **k):
        raise OllamaError("Ollama is down")
    monkeypatch.setattr(ollama_client, "generate", boom)
    assert run_chat("hi")[-1] == {"type": "error", "message": "Ollama is down"}


def test_context_is_wrapped_as_background(fake):
    seen = {}

    def handler(model, prompt, system):
        if model != config.ROUTER_MODEL:
            seen["prompt"] = prompt
        return "GENERAL_TASK" if model == config.ROUTER_MODEL else "hello"
    fake.handler = handler
    run_chat("hi there", context="[General]\nmachine: laptop")
    assert "for your reference only" in seen["prompt"] and "machine: laptop" in seen["prompt"]


# --- tools --------------------------------------------------------------------------------

def test_parse_tool_call_literals_only():
    assert tools.parse_tool_call('TOOL_CALL: git_log(n=3)') == ("git_log", [], {"n": 3})
    assert tools.parse_tool_call('x\nTOOL_CALL: web_search("cats")') == ("web_search", ["cats"], {})
    assert tools.parse_tool_call('TOOL_CALL: f(__import__("os").system("calc"))') is None
    assert tools.parse_tool_call("no call here") is None


def test_git_push_is_gated_and_not_model_callable():
    assert "git_push" not in tools.MODEL_TOOLS
    with pytest.raises(PermissionError):
        tools.git_push(".")
    # a model typing "git_push" only reaches the ask-the-user stub, which never pushes on its own
    assert tools.call_tool("git_push", [], {"confirmed": True}).startswith("Error")
    assert tools.call_tool("git_push", [], {}).startswith("Error: a push can only be requested from the chat")


def test_call_tool_never_raises():
    assert tools.call_tool("nope").startswith("Error: unknown tool")
    assert tools.call_tool("git_log", [], {"bogus": 1}).startswith("Error: bad arguments")


def test_directory_tree_ignores_junk(tmp_path):
    (tmp_path / "a.py").write_text("x")
    (tmp_path / ".git").mkdir()
    (tmp_path / "__pycache__").mkdir()
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.py").write_text("y")
    tree = tools.get_directory_tree(str(tmp_path))
    assert "a.py" in tree and "b.py" in tree and ".git" not in tree and "__pycache__" not in tree


def test_git_status_honest_outside_repo(tmp_path):
    assert "not a git repository" in tools.git_status(str(tmp_path)).lower()


# --- memory ---------------------------------------------------------------------------------------

@pytest.fixture
def mem(tmp_path):
    m = MemoryManager(tmp_path / "m.db")
    yield m
    m.graph.close()


def test_rolling_history_is_bounded_and_ordered(mem):
    chat = mem.create_chat()
    for i in range(12):
        mem.append_session_turn(chat["id"], f"question {i}", "GENERAL_TASK", f"answer {i}")
    history = mem.graph.get_node("session:" + chat["id"])["props"]["history"]
    assert len(history) == config.SESSION_HISTORY_TURNS
    assert "question 11" in history[-1] and "question 4" in history[0]


def test_context_tiers_and_cap(mem):
    mem.seed_general({"machine": "laptop"})
    proj = mem.create_project("Alpha", notes="x" * 3000)
    chat = mem.create_chat(project_key=proj["id"])
    mem.append_session_turn(chat["id"], "my name is Sam", "GENERAL_TASK", "hello Sam")
    ctx = mem.build_context(chat["id"])
    assert "[General]" in ctx and "[Project]" in ctx and "Sam" in ctx
    for block in ctx.split("\n\n["):
        assert len(block) <= config.MEMORY_CAP_CHARS + 40
    assert "· User" in ctx  # list rendered as bullet lines, not a Python repr
    assert "[" not in ctx.split("recent:")[1].split("\n")[1][:1]


def test_chats_do_not_leak_into_each_other(mem):
    a, b = mem.create_chat(), mem.create_chat()
    mem.append_session_turn(a["id"], "secret-alpha", "GENERAL_TASK", "ok")
    assert "secret-alpha" in mem.build_context(a["id"])
    assert "secret-alpha" not in mem.build_context(b["id"])


def test_long_memory_is_opt_in(mem):
    chat = mem.create_chat()
    mem.append_session_turn(chat["id"], "q", "GENERAL_TASK", "a")
    mem.add_long_memory("Summary of earlier work.", ["session:" + chat["id"]])
    assert "Summary of earlier work." not in mem.build_context(chat["id"])
    assert "Summary of earlier work." in mem.build_context(chat["id"], include_long_memory=True)


def test_chat_and_project_management(mem):
    p = mem.create_project("Proj")
    c = mem.create_chat(project_key=p["id"])
    mem.add_message(c["id"], "user", "hello")
    mem.add_message(c["id"], "assistant", "hi", {"category": "GENERAL_TASK"})
    assert [m["role"] for m in mem.get_messages(c["id"])] == ["user", "assistant"]
    assert mem.autotitle(c["id"], "How do I reverse a list in Python?")["title"].startswith("How do I reverse")
    assert mem.move_chat(c["id"], None)["project_key"] is None
    assert mem.move_chat(c["id"], p["id"])["project_key"] == p["id"]
    mem.delete_project(p["id"])  # chats survive, ungrouped
    assert mem.chat_summary(c["id"])["project_key"] is None
    p2 = mem.create_project("P2")
    c2 = mem.create_chat(project_key=p2["id"])
    mem.delete_project(p2["id"], delete_chats=True)
    with pytest.raises(KeyError):
        mem.chat_summary(c2["id"])
    mem.delete_chat(c["id"])
    assert mem.list_chats() == []


def test_persistence_across_restart(tmp_path):
    db = tmp_path / "p.db"
    m1 = MemoryManager(db)
    c = m1.create_chat(title="Kept")
    m1.add_message(c["id"], "user", "remember me")
    m1.graph.close()
    m2 = MemoryManager(db)
    assert m2.chat_summary(c["id"])["title"] == "Kept"
    assert m2.get_messages(c["id"])[0]["content"] == "remember me"
    m2.graph.close()


def test_gather_for_summary_empty_and_filled(mem):
    assert mem.gather_for_summary() == []
    c = mem.create_chat()
    mem.append_session_turn(c["id"], "q", "CODE_TASK", "a")
    assert len(mem.gather_for_summary()) == 1


def test_parse_tool_calls_handles_several_and_junk():
    text = 'TOOL_CALL: list_workspace()\nnoise\nTOOL_CALL: read_file(path="a.py")\nTOOL_CALL: broken(\nTOOL_CALL: __import__("os")'
    names = [c[0] for c in tools.parse_tool_calls(text)]
    assert names == ["list_workspace", "read_file", "__import__"]  # unknown names are refused later, at call time
    assert tools.call_tool("__import__", ["os"], registry={"list_workspace": lambda: ""}).startswith("Error: unknown tool")


def test_officekit_builds_every_format(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    import officekit as kit
    import pymupdf
    from openpyxl import load_workbook
    from pptx import Presentation
    chart = kit.make_chart("sales.png", "bar", ["Jan", "Feb"], {"2025": [1, 2], "2026": [2, 3]}, "Sales", "m", "u")
    for kind in ("line", "pie", "barh"):
        assert Path(kit.make_chart(f"{kind}.png", kind, ["a", "b"], [1, 2])).stat().st_size > 1000
    sections = [("Summary", "One.\n\nTwo & <b>three</b>."), ("Points", ["x", "y"]),
                ("Table", [["Item", "Cost"], ["Pens", 12]]), ("Chart", {"image": "sales.png", "caption": "c"}),
                ("Missing image", {"image": "nope.png"})]
    docx = kit.make_docx("out/r.docx", "Report", sections, subtitle="sub")
    text = tools.read_docx_text(docx)
    assert "Two & <b>three</b>." in text and "Pens | 12" in text and "Points" in text
    xlsx = load_workbook(kit.make_xlsx("out/d.xlsx", {"Costs": [["Item", "Cost"], ["Pens", 12]]}))
    assert xlsx["Costs"]["A2"].value == "Pens" and xlsx["Costs"]["A1"].font.bold
    assert load_workbook(kit.make_xlsx("out/single.xlsx", [["h"], [1]])).sheetnames == ["Sheet1"]
    deck = Presentation(kit.make_pptx("out/p.pptx", "Deck", [("A", ["b1", "b2"]), ("Chart", {"image": chart}), ("T", "txt")]))
    assert len(deck.slides) == 4
    pdf = pymupdf.open(kit.make_pdf("out/r.pdf", "Report", sections))
    assert "Report" in pdf[0].get_text() and "Pens" in pdf[0].get_text() and Path("out/r.pdf").stat().st_size < 400_000
    pdf.close()
    assert Path(kit.make_csv("out/c.csv", [["a", "b"], [1, "é"]])).read_text(encoding="utf-8-sig").splitlines()[1] == "1,é"


# --- API (through the real FastAPI stack) ---------------------------------------------------------------

@pytest.fixture
def client(fake):
    from fastapi.testclient import TestClient
    import main
    with TestClient(main.app, base_url="http://127.0.0.1:8420") as c:
        yield c


def sse(response):
    return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]


def test_api_chat_roundtrip_and_memory(client, fake):
    fake.handler = lambda m, p, s: "GENERAL_TASK" if m == config.ROUTER_MODEL else "Nice to meet you, Sam."
    chat = client.post("/api/chats", json={}).json()
    r1 = client.post("/api/chat", json={"message": "My name is Sam", "chat_id": chat["id"]})
    assert r1.status_code == 200 and r1.json()["reply"].startswith("Nice to meet")
    assert r1.json()["chat"]["title"] == "My name is Sam"
    seen = {}

    def spy(m, p, s):
        if m != config.ROUTER_MODEL:
            seen["p"] = p
        return "GENERAL_TASK" if m == config.ROUTER_MODEL else "Your name is Sam."
    fake.handler = spy
    client.post("/api/chat", json={"message": "What is my name?", "chat_id": chat["id"]})
    assert "My name is Sam" in seen["p"]  # second turn sees the first
    msgs = client.get(f"/api/chats/{chat['id']}/messages").json()["messages"]
    assert [m["role"] for m in msgs] == ["user", "assistant", "user", "assistant"]


def test_api_stream_emits_phases_tokens_done(client, fake):
    fake.handler = lambda m, p, s: "MATH_TASK" if m == config.ROUTER_MODEL else "It is 4."
    events = sse(client.post("/api/chat/stream", json={"message": "2+2?"}))
    kinds = [e["type"] for e in events]
    assert kinds[0] == "chat" and "token" in kinds and kinds[-1] == "done"
    assert [e["phase"] for e in events if e["type"] == "phase"][:2] == ["classifying", "routed"]
    assert events[-1]["chat"]["title"] == "2+2?"


def test_api_stream_error_is_reported_and_saved(client, fake, monkeypatch):
    def boom(*a, **k):
        raise OllamaError("Ollama is down")
    monkeypatch.setattr(ollama_client, "generate", boom)
    chat = client.post("/api/chats", json={}).json()
    events = sse(client.post("/api/chat/stream", json={"message": "hi", "chat_id": chat["id"]}))
    assert events[-1] == {"type": "error", "message": "Ollama is down"}
    assert client.get(f"/api/chats/{chat['id']}/messages").json()["messages"][-1]["role"] == "error"
    assert client.post("/api/chat", json={"message": "hi", "chat_id": chat["id"]}).status_code == 503


def test_api_chats_and_projects_crud(client):
    p = client.post("/api/projects", json={"name": "Alpha"}).json()
    c = client.post("/api/chats", json={"project_key": p["id"]}).json()
    listing = client.get("/api/chats").json()
    assert any(x["id"] == c["id"] and x["project_key"] == p["id"] for x in listing["chats"])
    assert client.patch(f"/api/chats/{c['id']}", json={"title": "Renamed"}).json()["title"] == "Renamed"
    assert client.post(f"/api/chats/{c['id']}/move", json={"project_key": None}).json()["project_key"] is None
    assert client.patch(f"/api/projects/{p['id']}", json={"name": "Beta"}).json()["name"] == "Beta"
    assert client.delete(f"/api/chats/{c['id']}").status_code == 200
    assert client.delete(f"/api/chats/{c['id']}").status_code == 404
    assert client.post("/api/chats", json={"project_key": "nope"}).status_code == 404
    assert client.delete(f"/api/projects/{p['id']}").status_code == 200


def test_api_git_push_needs_confirmation(client):
    assert client.post("/api/git/push", json={"repo_path": "."}).status_code == 403
    assert client.post("/api/git/push", json={"repo_path": ".", "confirmed": False}).status_code == 403


def test_api_status_repo_tree_memory_and_ui(client):
    s = client.get("/api/status").json()
    assert s["threads"] == config.NUM_THREAD and any(m["role"] == "DEEP" for m in s["models"])
    assert "config.py" in client.get("/api/repo-tree", params={"path": str(ROOT)}).json()["output"]
    assert "context" in client.get("/api/memory").json()
    assert "Local AI" in client.get("/").text


def test_deep_mode_never_runs_tools(fake):
    fake.handler = lambda m, p, s: 'TOOL_CALL: extract_pdf_text(path="intervals.pdf")'
    done = run_chat("write a function", deep=True)[-1]
    assert done["meta"]["tool"] is None and "TOOL_CALL" in done["reply"]
    assert fake.calls == [config.ESCALATION_MODEL]  # one generation, no follow-up


def test_push_review_returns_only_the_git_error_outside_a_repo(tmp_path):
    out = tools.git_diff_for_push_review(str(tmp_path))
    assert out.lower().startswith("fatal:") and "usage:" not in out


def test_ensure_running(monkeypatch):
    monkeypatch.setattr(ollama_client, "list_models", lambda: ["x"])
    assert ollama_client.ensure_running(wait=1) is True

    def down():
        raise OllamaError("down")
    monkeypatch.setattr(ollama_client, "list_models", down)
    monkeypatch.setenv("LOCALAPPDATA", str(ROOT / "no-such-dir"))
    assert ollama_client.ensure_running(wait=1) is False  # not installed: gives up quietly


def test_api_web_search_and_git_helpers(client, monkeypatch, tmp_path):
    monkeypatch.setattr(tools, "web_search_results",
                        lambda q, n=5: [{"title": "T", "href": "https://example.com", "body": "B"}])
    assert client.post("/api/web-search", json={"query": "x"}).json()["results"][0]["title"] == "T"

    def offline(q, n=5):
        raise RuntimeError("offline")
    monkeypatch.setattr(tools, "web_search_results", offline)
    assert client.post("/api/web-search", json={"query": "x"}).status_code == 502
    for kind in ("log", "diff", "status"):
        r = client.post(f"/api/git/{kind}", json={"repo_path": str(tmp_path)})
        assert r.status_code == 200 and "not a git repository" in r.json()["output"].lower()
    assert client.get("/api/status").json()["project_dir"] == str(config.BASE_DIR)


# --- documents and new tools ----------------------------------------------------------------------

def make_pdf(*page_texts) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    for text in page_texts:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(*paragraphs, table=None) -> bytes:
    import io
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_pdf_docx_and_text():
    pdf = tools.extract_document("a.pdf", make_pdf("Hello page one", "Second page"))
    assert pdf["kind"] == "pdf" and pdf["pages"] == 2 and "Hello page one" in pdf["text"] and "Second page" in pdf["text"]
    docx = tools.extract_document("b.DOCX", make_docx("First para", "Second para", table=[["a", "b"], ["c", "d"]]))
    assert docx["kind"] == "docx" and "First para" in docx["text"] and "a | b" in docx["text"]
    txt = tools.extract_document("c.py", "print('héllo')\n".encode("utf-8"))
    assert txt["kind"] == "text" and "héllo" in txt["text"]


def test_extract_rejects_bad_files():
    with pytest.raises(ValueError):
        tools.extract_document("virus.exe", b"MZ")
    with pytest.raises(ValueError):
        tools.extract_document("fake.txt", b"abc\x00def")
    assert "scanned" in tools.extract_document("blank.pdf", make_pdf(""))["text"]


def test_path_based_pdf_and_docx_tools(tmp_path):
    (tmp_path / "r.pdf").write_bytes(make_pdf("Quarterly report"))
    (tmp_path / "n.docx").write_bytes(make_docx("Meeting notes"))
    assert "Quarterly report" in tools.extract_pdf_text(str(tmp_path / "r.pdf"))
    assert "Meeting notes" in tools.read_docx_text(str(tmp_path / "n.docx"))
    assert tools.extract_pdf_text(str(tmp_path / "missing.pdf")).startswith("Error: file not found")
    out = tools.pdf_to_docx(str(tmp_path / "r.pdf"), str(tmp_path / "out" / "r.docx"))
    assert out.startswith("Wrote") and "Quarterly report" in tools.read_docx_text(str(tmp_path / "out" / "r.docx"))


def test_read_file_refuses_secrets_and_binaries(tmp_path):
    (tmp_path / "ok.py").write_text("x = 1\n" * 3000)
    (tmp_path / ".env").write_text("TOKEN=abc")
    (tmp_path / "id_rsa").write_text("-----BEGIN-----")
    (tmp_path / "ssh").mkdir()
    (tmp_path / ".ssh").mkdir()
    (tmp_path / ".ssh" / "config").write_text("Host x")
    (tmp_path / "blob.bin").write_bytes(b"\x00\x01\x02")
    assert "truncated" in tools.read_file(str(tmp_path / "ok.py"))
    for secret in (".env", "id_rsa", ".ssh/config"):
        assert "refusing" in tools.read_file(str(tmp_path / secret))
    assert "binary" in tools.read_file(str(tmp_path / "blob.bin"))
    assert "refusing" in tools.read_file(str(config.DB_PATH))


def test_search_files(tmp_path):
    (tmp_path / "a.py").write_text("def foo():\n    # TODO: fix\n    pass\n")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.md").write_text("nothing\ntodo later\n")
    (tmp_path / "node_modules").mkdir()
    (tmp_path / "node_modules" / "c.js").write_text("// TODO hidden")
    (tmp_path / ".env").write_text("TODO=secret")
    out = tools.search_files("todo", str(tmp_path))
    assert "a.py:2" in out and "sub/b.md:2" in out and "node_modules" not in out and ".env" not in out
    assert tools.search_files("zzz-none", str(tmp_path)).startswith("No matches")
    assert tools.search_files("x", str(tmp_path / "missing")).startswith("Error")


def test_html_to_text_and_fetch_safety(monkeypatch):
    import httpx
    html = ("<html><head><title>Docs</title><style>p{}</style></head><body><nav>menu</nav>"
            "<h1>Intro</h1><p>Hello <b>world</b></p><script>alert(1)</script><footer>foot</footer></body></html>")
    text = tools.html_to_text(html)
    assert text.startswith("Docs") and "Hello world" in text and "alert" not in text and "menu" not in text

    for bad in ("http://127.0.0.1:8420/api/status", "http://localhost/x", "file:///etc/passwd", "ftp://example.com"):
        assert tools.fetch_webpage(bad).startswith("Error:"), bad

    monkeypatch.setattr(tools, "_assert_public_host", lambda host: None)
    real_client = httpx.Client

    def client(**kw):
        def handler(request):
            if request.url.path == "/old":
                return httpx.Response(302, headers={"location": "/new"})
            return httpx.Response(200, headers={"content-type": "text/html"}, text=html)
        return real_client(transport=httpx.MockTransport(handler), **kw)
    monkeypatch.setattr(httpx, "Client", client)
    assert "Hello world" in tools.fetch_webpage("https://example.com/old")  # follows the redirect
    assert "Hello world" in tools.fetch_webpage_text("https://example.com/new", max_chars=5000)


def test_system_info_and_registry():
    info = tools.get_system_info()
    assert "RAM:" in info and "CPU:" in info and "Disk" in info
    for name in ("extract_pdf_text", "read_docx_text", "read_file", "search_files", "fetch_webpage", "get_system_info"):
        assert name in tools.MODEL_TOOLS
    assert "git_push" not in tools.MODEL_TOOLS and "write_file" not in tools.MODEL_TOOLS


@pytest.mark.parametrize("message", [
    r"Extract the text from C:\Users\rayya\invoice.pdf and tell me the total",
    "Fetch https://example.com and tell me what it says",
    "Search my project for every TODO comment",
    "search the folder for the word invoice",
    "How much free RAM and disk space does this computer have?",
    "How much RAM do I have right now?",
    r"Read the file C:\Users\sam\notes\draft.docx",
    "what's the git status",
    "show me the files in my project folder",
    "Find every TODO in my project",
    "Summarize https://docs.python.org/3/ for me",
])
def test_rules_send_tool_requests_to_the_coder_without_the_router(fake, message):
    assert router_engine.rule_category(message) == "CODE_TASK"
    fake.handler = lambda m, p, s: "ok"
    events = run_chat(message)
    assert fake.calls == [config.SPECIALISTS["CODE_TASK"]]  # the router model was never called
    assert events[-1]["meta"]["category"] == "CODE_TASK"


@pytest.mark.parametrize("message", [
    "Write a polite email asking my landlord to fix the heating",
    "What is 15% of 2400 plus 87?",
    "How do I open port 8080 in the Windows firewall?",
    "Who painted the Mona Lisa?",
    "Write an email that links to https://example.com for the invoice",
    "Find a good name for my project",
    "Hi! My name is Sam and I'm building a local AI system on my laptop.",
    "How do I check disk usage on Linux?",
    "--- Attached: report.pdf ---\nSee https://x.com\n--- end of report.pdf ---\n\nSummarize this",
    "Explain what a PDF is",
])
def test_rules_leave_everything_else_to_the_router(message):
    assert router_engine.rule_category(message) is None


def test_classifier_input_is_bounded(fake):
    seen = {}

    def handler(model, prompt, system):
        if model == config.ROUTER_MODEL:
            seen["prompt"] = prompt
            return "WRITING_TASK"
        return "ok"
    fake.handler = handler
    run_chat("Summarize this:\n" + "x" * 20000 + "\nthanks")
    assert len(seen["prompt"]) < len(router_engine.CLASSIFIER_PROMPT) + 1200
    assert seen["prompt"].rstrip().endswith("Label:") and "thanks" in seen["prompt"]


def test_api_documents(client, tmp_path):
    pdf = make_pdf("Alpha beta gamma")
    r = client.post("/api/docs/extract", params={"filename": "x.pdf"}, content=pdf)
    assert r.status_code == 200 and "Alpha beta gamma" in r.json()["text"] and r.json()["pages"] == 1
    assert client.post("/api/docs/extract", params={"filename": "x.exe"}, content=b"MZ").status_code == 415
    assert client.post("/api/docs/extract", params={"filename": "x.pdf"}, content=b"").status_code == 400
    assert client.post("/api/docs/extract", params={"filename": "x.pdf"}, content=b"not a pdf").status_code == 422
    conv = client.post("/api/docs/pdf-to-docx", params={"filename": "x.pdf"}, content=pdf)
    assert conv.status_code == 200 and "Alpha beta gamma" in tools.docx_text_from_bytes(conv.content)
    assert client.post("/api/docs/pdf-to-docx", content=make_pdf("")).status_code == 415


def test_api_search_fetch_system(client, tmp_path):
    (tmp_path / "m.py").write_text("needle = 1\n")
    out = client.post("/api/files/search", json={"query": "needle", "path": str(tmp_path)}).json()["output"]
    assert "m.py:1" in out
    assert client.post("/api/web/fetch", json={"url": "http://127.0.0.1:8420/"}).status_code == 400
    assert "RAM:" in client.get("/api/system").json()["output"]


def test_api_rejects_foreign_host_header(fake):
    from fastapi.testclient import TestClient
    import main
    with TestClient(main.app, base_url="http://evil.example.com") as c:
        assert c.get("/api/status").status_code == 400
