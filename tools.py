"""Tools: docs, search, git, quality checks, files.

`MODEL_TOOLS` is the only table models can reach. `git_push` is deliberately
absent from it and is hard-gated in code (`confirmed=True` required).
"""
import ast
import inspect
import io
import ipaddress
import os
import platform
import re
import socket
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse

import config
import constitution

CREATE_NO_WINDOW = 0x08000000 if sys.platform == "win32" else 0
MAX_UPLOAD_BYTES = 30 * 1024 * 1024
TEXT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".py", ".json", ".csv", ".tsv", ".log", ".js", ".ts", ".tsx", ".jsx",
    ".html", ".htm", ".css", ".yml", ".yaml", ".toml", ".ini", ".cfg", ".xml", ".sql", ".sh",
    ".bat", ".ps1", ".c", ".h", ".cpp", ".hpp", ".java", ".go", ".rs", ".rb", ".php", ".tex",
}


def _resolve(path) -> Path:
    """Relative paths are relative to the project folder, not the server's cwd."""
    p = Path(str(path)).expanduser()
    return p if p.is_absolute() else config.BASE_DIR / p


# --- Documents ------------------------------------------------------------

NO_TEXT_PDF = "(no extractable text; this looks like a scanned PDF and OCR is not supported)"


def pdf_text_from_bytes(data: bytes, max_pages: int = 500) -> tuple:
    """(text, total_pages) from PDF bytes."""
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        total = doc.page_count
        pages = [page.get_text() for page in doc if page.number < max_pages]
    return "\n".join(pages).strip(), total


def docx_text_from_bytes(data: bytes) -> str:
    """Paragraphs and table rows of a .docx, one per line."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())


def pdf_bytes_to_docx_bytes(data: bytes) -> bytes:
    """Text-only PDF -> Word conversion, in memory."""
    from docx import Document

    text, _ = pdf_text_from_bytes(data, max_pages=10_000)
    if not text:
        raise ValueError(NO_TEXT_PDF.strip("()"))
    doc = Document()
    for block in re.split(r"\n{2,}", text):
        doc.add_paragraph(block)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def extract_document(filename: str, data: bytes) -> dict:
    """Extract text from an uploaded PDF, Word or plain-text/code file.

    Returns {"name", "kind", "text", "pages", "chars"}. Raises ValueError for
    unsupported types.
    """
    ext = Path(filename or "").suffix.lower()
    pages = None
    if ext == ".pdf":
        kind = "pdf"
        text, pages = pdf_text_from_bytes(data)
        text = text or NO_TEXT_PDF
    elif ext == ".docx":
        kind = "docx"
        text = docx_text_from_bytes(data) or "(this Word file has no text)"
    elif ext in TEXT_EXTENSIONS:
        kind = "text"
        if b"\x00" in data[:4096]:
            raise ValueError(f"{filename} looks like a binary file, not text.")
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type '{ext or filename}'. Use PDF, DOCX, or a text/code file.")
    return {"name": filename, "kind": kind, "text": text, "pages": pages, "chars": len(text)}


def extract_pdf_text(path: str, max_pages: int = 50) -> str:
    """Extract the text of a PDF file at the given path."""
    p = _resolve(path)
    if not p.is_file():
        return f"Error: file not found: {p}"
    text, _ = pdf_text_from_bytes(p.read_bytes(), max_pages=int(max_pages))
    return text or NO_TEXT_PDF


def read_docx_text(path: str) -> str:
    """Extract the text of a Word (.docx) file at the given path."""
    p = _resolve(path)
    if not p.is_file():
        return f"Error: file not found: {p}"
    return docx_text_from_bytes(p.read_bytes()) or "(this Word file has no text)"


def pdf_to_docx(pdf_path: str, docx_path: str) -> str:
    """Copy a PDF's text into a new .docx (text only, no layout)."""
    src = _resolve(pdf_path)
    if not src.is_file():
        return f"Error: file not found: {src}"
    out = _resolve(docx_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(pdf_bytes_to_docx_bytes(src.read_bytes()))
    return f"Wrote {out}"


# --- Web search -----------------------------------------------------------

def _refuse_secret(what: str, text: str) -> None:
    """Article 1.3: nothing that looks like a secret may leave this computer in a search or URL."""
    kinds = constitution.find_secrets(text)
    if kinds:
        constitution.audit("outbound_blocked", what=what, kind=kinds[0])
        raise ValueError(f"Refused by the constitution (data security): the {what} contains something that "
                         f"looks like a secret ({kinds[0]}). It was not sent.")


def web_search_results(query: str, max_results: int = 5) -> list:
    """Structured search results: [{title, href, body}]. Raises on network trouble."""
    _refuse_secret("search query", query)
    try:
        from ddgs import DDGS
    except ImportError:
        from duckduckgo_search import DDGS
    raw = DDGS().text(query, max_results=int(max_results))
    return [{"title": r.get("title", ""), "href": r.get("href", ""), "body": r.get("body", "")}
            for r in raw]


def web_search(query: str, max_results: int = 5) -> str:
    """Search the web (DuckDuckGo, no key). Returns titles, links and snippets."""
    try:
        results = web_search_results(query, max_results)
    except Exception as exc:  # network, rate limit, missing package...
        return f"Web search failed: {type(exc).__name__}: {exc}"
    if not results:
        return "No results."
    return "\n".join(f"{i}. {r['title']}\n   {r['href']}\n   {r['body']}"
                     for i, r in enumerate(results, 1))


class _TextExtractor(HTMLParser):
    """Visible text of an HTML page; skips scripts, styles and page chrome."""

    SKIP = {"script", "style", "noscript", "svg", "head", "nav", "footer", "form", "iframe"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "section", "article"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self._skip += 1
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _TextExtractor()
    parser.feed(html)
    body = re.sub(r"[ \t\r\f\v]+", " ", "".join(parser.parts))
    body = re.sub(r"\n\s*\n+", "\n\n", body).strip()
    title = parser.title.strip()
    return f"{title}\n\n{body}" if title else body


def _assert_public_host(host: str) -> None:
    """Refuse loopback / private / link-local targets so the model can't probe the LAN."""
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror as exc:
        raise ValueError(f"Cannot resolve {host}: {exc}") from None
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if not ip.is_global:
            raise ValueError(f"Refusing to fetch {host}: it points to a private or local address.")


def fetch_webpage_text(url: str, max_chars: int = 6000) -> str:
    """Download a public web page and return its readable text. Raises ValueError for refusals."""
    import httpx

    _refuse_secret("web address", url)
    current = url.strip()
    with httpx.Client(timeout=20.0, follow_redirects=False, headers={"User-Agent": "LocalAI/1.0"}) as client:
        for _ in range(4):
            parsed = urlparse(current)
            if parsed.scheme not in ("http", "https") or not parsed.hostname:
                raise ValueError("Only http:// and https:// URLs can be fetched.")
            _assert_public_host(parsed.hostname)
            try:
                resp = client.get(current)
            except httpx.HTTPError as exc:
                raise ValueError(f"Could not fetch the page: {type(exc).__name__}: {exc}") from None
            if resp.is_redirect and resp.headers.get("location"):
                current = urljoin(current, resp.headers["location"])
                continue
            break
        else:
            raise ValueError("Too many redirects.")
    if resp.status_code >= 400:
        raise ValueError(f"The server answered HTTP {resp.status_code}.")
    ctype = resp.headers.get("content-type", "")
    if "html" in ctype or "xml" in ctype or not ctype:
        text = html_to_text(resp.text[:2_000_000])
    elif ctype.startswith("text/") or "json" in ctype:
        text = resp.text[:2_000_000]
    else:
        raise ValueError(f"Not a text page (content type {ctype or 'unknown'}).")
    limit = int(max_chars)
    return text[:limit] + (f"\n...[truncated, {len(text) - limit} more characters]" if len(text) > limit else "")


def fetch_webpage(url: str, max_chars: int = 6000) -> str:
    """Download a public web page (http/https) and return its readable text."""
    try:
        return fetch_webpage_text(url, max_chars)
    except ValueError as exc:
        return f"Error: {exc}"


# --- Git ------------------------------------------------------------------

def _git(repo_path, *args, timeout: int = 60) -> str:
    repo = _resolve(repo_path or ".")
    env = dict(os.environ, GIT_TERMINAL_PROMPT="0")
    try:
        proc = subprocess.run(
            ["git", "-C", str(repo), *args],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=timeout, env=env, stdin=subprocess.DEVNULL,
            creationflags=CREATE_NO_WINDOW,
        )
    except FileNotFoundError:
        return "Error: git is not installed or not on PATH."
    except subprocess.TimeoutExpired:
        return f"Error: git {args[0]} timed out after {timeout}s."
    out = (proc.stdout + proc.stderr).strip()
    return out or "(no output)"


def git_status(repo_path: str = ".") -> str:
    """Show working tree status of a git repository."""
    return _git(repo_path, "status", "--short", "--branch")


def git_pull(repo_path: str = ".") -> str:
    """Pull the latest changes (fast-forward only)."""
    return _git(repo_path, "pull", "--ff-only", timeout=120)


def git_fetch(repo_path: str = ".") -> str:
    return _git(repo_path, "fetch", timeout=120)


def git_log(repo_path: str = ".", n: int = 10) -> str:
    """Show the last n commits, one per line."""
    return _git(repo_path, "log", "--oneline", "--decorate", f"-{int(n)}")


def git_diff(repo_path: str = ".") -> str:
    """Show uncommitted changes."""
    return _git(repo_path, "diff")


def git_branch(repo_path: str = ".") -> str:
    return _git(repo_path, "branch", "--all")


def _git_run(repo_path, *args, timeout: int = 60) -> tuple:
    """(returncode, stdout, stderr) of a git command in the repo."""
    repo = _resolve(repo_path or ".")
    env = dict(os.environ, GIT_TERMINAL_PROMPT="0")
    try:
        proc = subprocess.run(
            ["git", "-C", str(repo), *args],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=timeout, env=env, stdin=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW,
        )
    except FileNotFoundError:
        return 127, "", "git is not installed or not on PATH."
    except subprocess.TimeoutExpired:
        return 124, "", f"git {args[0]} timed out after {timeout}s."
    return proc.returncode, proc.stdout, proc.stderr


_SAFE_REF = re.compile(r"^[A-Za-z0-9_][A-Za-z0-9_./-]*$")


def _safe_ref(name: str, what: str) -> str:
    """Branch/remote names go into git argument lists: refuse anything that could be an option."""
    if not isinstance(name, str) or not _SAFE_REF.match(name) or ".." in name or len(name) > 200:
        raise ValueError(f"Invalid {what} name: {name!r}")
    return name


def _repo_text(path: Path, limit: int = 1_000_000):
    """Text of a small text file, or None."""
    try:
        if not path.is_file() or path.stat().st_size > limit:
            return None
        data = path.read_bytes()
    except OSError:
        return None
    return None if b"\x00" in data[:4096] else data.decode("utf-8", errors="replace")


def git_add(repo_path: str = ".", paths=".") -> str:
    """Stage files for commit. Refuses secrets and sensitive files (data security)."""
    items = [paths] if isinstance(paths, str) else [str(p) for p in paths]
    if any(i.startswith("-") for i in items):
        return "Error: paths may not start with '-'."
    rc, out, err = _git_run(repo_path, "add", "--dry-run", "--", *items)
    if rc != 0:
        return (err or out).strip() or "Error: git add failed."
    candidates = [m.group(1) for line in out.splitlines()
                  if (m := re.match(r"^(?:add|remove) '(.+)'$", line.strip()))]
    repo = _resolve(repo_path or ".")
    sensitive = [c for c in candidates if constitution.is_sensitive_path(c)]
    leaks = []
    for c in candidates[:300]:
        text = None if c in sensitive else _repo_text(repo / c)
        if text:
            leaks.extend(f"{c}: {kind}" for kind in constitution.find_secrets(text))
    if sensitive or leaks:
        constitution.audit("stage_blocked", repo=str(repo), sensitive=sensitive, secrets=leaks)
        lines = ["Error: refused by the constitution (data security). Nothing was staged."]
        lines += [f"  sensitive file: {s}" for s in sensitive] + [f"  looks like a secret: {s}" for s in leaks]
        lines.append("Remove or git-ignore these files, then try again. Only the user can override this.")
        return "\n".join(lines)
    result = _git(repo_path, "add", "--", *items)
    return f"Staged {len(candidates)} file(s)." + ("" if result == "(no output)" else "\n" + result)


def git_commit(repo_path: str, message: str) -> str:
    """Commit the staged changes. Refuses when the staged changes contain secrets."""
    if not message or not message.strip():
        return "Error: commit message is required."
    rc, names, _ = _git_run(repo_path, "diff", "--cached", "--name-only")
    sensitive = [n for n in names.splitlines() if constitution.is_sensitive_path(n)]
    rc, patch, _ = _git_run(repo_path, "diff", "--cached", "-U0", "--no-color")
    leaks = [f"{f}: {kind}" for f, kind in constitution.scan_patch(patch)]
    if sensitive or leaks:
        constitution.audit("commit_blocked", sensitive=sensitive, secrets=leaks)
        return ("Error: refused by the constitution (data security). The staged changes contain "
                + "; ".join(sensitive + leaks) + ". Unstage them first.")
    result = _git(repo_path, "commit", "-m", message[:2000])
    if "Author identity unknown" in result or "unable to auto-detect email" in result:
        return ("Error: git has no author name/email configured on this computer, so nothing was committed. "
                "Only the user can set it, once: git config --global user.name \"Your Name\" and "
                "git config --global user.email \"you@example.com\". Tell the user this and stop; do not try to work around it.")
    return result


def _remote_host(url: str) -> str:
    """Host of a git remote URL; 'local' for folder remotes."""
    u = url.strip()
    m = re.match(r"^[\w.-]+@([^:/]+):", u)
    if m:
        return m.group(1).lower()
    if "://" in u:
        parsed = urlparse(u)
        return "local" if parsed.scheme == "file" else (parsed.hostname or "local").lower()
    return "local"


def git_push_review(repo_path: str = ".", remote: str = "origin", branch: str = "", force: bool = False) -> dict:
    """Everything the user needs to decide on a push, plus the constitution's verdict.

    Returns {"ok", "error", "repo", "remote", "branch", "remote_url", "host", "commits", "files",
    "nothing_to_push", "violations", "text"}. Never pushes.
    """
    result = {"ok": False, "error": "", "repo": str(_resolve(repo_path or ".")), "remote": remote, "branch": branch,
              "remote_url": "", "host": "", "commits": [], "files": [], "nothing_to_push": False,
              "violations": [], "text": ""}
    try:
        _safe_ref(remote, "remote")
        if branch:
            _safe_ref(branch, "branch")
    except ValueError as exc:
        result["error"] = str(exc)
        return result
    rc, out, err = _git_run(repo_path, "rev-parse", "--is-inside-work-tree")
    if rc != 0 or out.strip() != "true":
        result["error"] = (err or out).strip() or "Not a git repository."
        return result
    if not branch:
        rc, out, err = _git_run(repo_path, "rev-parse", "--abbrev-ref", "HEAD")
        branch = out.strip()
        if rc != 0 or branch in ("", "HEAD"):
            result["error"] = "No branch is checked out (detached HEAD, or no commits yet)."
            return result
    result["branch"] = branch
    rc, out, err = _git_run(repo_path, "remote", "get-url", remote)
    if rc != 0:
        result["error"] = f"There is no remote named '{remote}'. Add one with: git remote add {remote} <url>"
        return result
    url = out.strip()
    result["remote_url"], result["host"] = constitution.redact(url), _remote_host(url)

    rng = [branch, "--not", f"--remotes={remote}"]
    rc, out, err = _git_run(repo_path, "log", "--oneline", *rng)
    if rc != 0:
        result["error"] = (err or out).strip()
        return result
    commits = out.splitlines()
    result["commits"] = commits
    result["ok"] = True
    files, sizes, secrets = [], {}, []
    if commits:
        _, out, _ = _git_run(repo_path, "log", "--name-only", "--format=", *rng)
        files = list(dict.fromkeys(line.strip() for line in out.splitlines() if line.strip()))
        _, patch, _ = _git_run(repo_path, "log", "-p", "-U0", "--no-color", "--format=", *rng, timeout=120)
        secrets = constitution.scan_patch(patch[:5_000_000])
        for f in files[:400]:
            rc, size, _ = _git_run(repo_path, "cat-file", "-s", f"{branch}:{f}")
            if rc == 0 and size.strip().isdigit():
                sizes[f] = int(size.strip())
    else:
        result["nothing_to_push"] = True
    result["files"] = files
    result["violations"] = constitution.evaluate_push({
        "branch": branch, "host": result["host"], "force": force, "files": files, "sizes": sizes, "secrets": secrets})

    lines = [f"Branch: {branch}", f"Remote: {remote} ({result['remote_url']}, host {result['host']})"]
    if commits:
        lines += [f"Commits to push ({len(commits)}):"] + [f"  {c}" for c in commits[:30]]
        lines += [f"Files changed ({len(files)}):"] + [f"  {f}" for f in files[:60]]
    else:
        lines.append("Nothing to push: the remote already has every commit on this branch.")
    result["text"] = "\n".join(lines)
    return result


def git_diff_for_push_review(repo_path: str = ".", remote: str = "origin", branch: str = "") -> str:
    """What a push would send: commits, files and any constitution violations, as text."""
    review = git_push_review(repo_path, remote, branch)
    if not review["ok"]:
        error = review["error"]
        return error if error.lower().startswith(("fatal:", "error:")) else f"fatal: {error}"
    text = review["text"]
    for v in review["violations"]:
        text += f"\n[{v['tier'].upper()} RULE: {v['title']}] {v['message']}" + "".join(f"\n    {d}" for d in v["details"])
    return text


class PolicyBlocked(Exception):
    """A push the constitution forbids and the user has not overridden."""

    def __init__(self, violations: list) -> None:
        super().__init__("; ".join(f"{v['title']}: {v['message']}" for v in violations))
        self.violations = violations


def git_push(repo_path: str = ".", confirmed: bool = False, remote: str = "origin", branch: str = "",
             force_with_lease: bool = False) -> str:
    """Low-level push. Refuses unless the user confirmed; policy is enforced by guarded_push."""
    if confirmed is not True:
        raise PermissionError("git push requires explicit user confirmation (confirmed=True).")
    _safe_ref(remote, "remote")
    if not branch:
        _, out, _ = _git_run(repo_path, "rev-parse", "--abbrev-ref", "HEAD")
        branch = out.strip()
    _safe_ref(branch, "branch")
    rc, _, _ = _git_run(repo_path, "rev-parse", "--abbrev-ref", f"{branch}@{{u}}")
    args = ["push"] + (["--force-with-lease"] if force_with_lease else []) + (["--set-upstream"] if rc != 0 else [])
    return _git(repo_path, *args, remote, branch, timeout=180)


def guarded_push(repo_path: str = ".", remote: str = "origin", branch: str = "", force: bool = False,
                 overrides=(), confirmed: bool = False) -> dict:
    """The one place pushes happen. Needs the user's confirmation AND a clean constitution check.

    `overrides` are the rule ids the user explicitly overrode for this single push.
    Raises PermissionError (not confirmed), ValueError (cannot push), PolicyBlocked (rule violated).
    """
    if confirmed is not True:
        raise PermissionError("git push requires explicit user confirmation (confirmed=True).")
    review = git_push_review(repo_path, remote, branch, force)
    if not review["ok"]:
        raise ValueError(review["error"])
    if review["nothing_to_push"] and not force:
        return {"output": "Nothing to push: the remote already has these commits.", "review": review, "pushed": False}
    blocked = constitution.unresolved(review["violations"], overrides)
    if blocked:
        constitution.audit("push_blocked", repo=review["repo"], remote=review["remote"], branch=review["branch"],
                           rules=[v["rule"] for v in blocked])
        raise PolicyBlocked(blocked)
    used = [v["rule"] for v in review["violations"]]
    output = git_push(repo_path, True, remote, review["branch"], force_with_lease=force)
    constitution.audit("push_done", repo=review["repo"], remote=review["remote"], host=review["host"],
                       branch=review["branch"], commits=len(review["commits"]), overrides_used=used,
                       result=output[-300:])
    return {"output": output, "review": review, "pushed": True}


# --- Quality --------------------------------------------------------------

def run_ruff_check(path: str = ".") -> str:
    """Lint Python code with ruff."""
    target = _resolve(path)
    try:
        proc = subprocess.run(
            [sys.executable, "-m", "ruff", "check", str(target)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=120, stdin=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW,
        )
    except subprocess.TimeoutExpired:
        return "Error: ruff timed out."
    return (proc.stdout + proc.stderr).strip() or "(no output)"


def run_pytest(path: str = ".") -> str:
    """Run the pytest suite found at path."""
    target = _resolve(path)
    try:
        proc = subprocess.run(
            [sys.executable, "-m", "pytest", "-q", "--no-header", "-p", "no:cacheprovider", str(target)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=300, stdin=subprocess.DEVNULL, cwd=str(target if target.is_dir() else target.parent),
            creationflags=CREATE_NO_WINDOW,
        )
    except subprocess.TimeoutExpired:
        return "Error: pytest timed out after 300s."
    return (proc.stdout + proc.stderr).strip() or "(no output)"


def run_python_file(path, timeout: int = config.AGENT_RUN_TIMEOUT, cwd=None, env=None) -> dict:
    """Run a Python script; stdin is closed so input() fails fast instead of hanging."""
    target = _resolve(path)
    try:
        proc = subprocess.run(
            [sys.executable, str(target)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=timeout, stdin=subprocess.DEVNULL, env=env,
            cwd=str(cwd or target.parent), creationflags=CREATE_NO_WINDOW,
        )
    except subprocess.TimeoutExpired as exc:
        out = exc.stdout if isinstance(exc.stdout, str) else ""
        return {"returncode": -1, "stdout": out, "stderr": f"Timed out after {timeout}s.", "timed_out": True}
    return {"returncode": proc.returncode, "stdout": proc.stdout, "stderr": proc.stderr, "timed_out": False}


# --- Files ----------------------------------------------------------------

IGNORED_DIRS = {
    ".git", "node_modules", "__pycache__", ".venv", "venv", ".idea", ".vscode",
    ".pytest_cache", ".ruff_cache", ".mypy_cache", "dist", "build",
}


def _is_sensitive(p: Path) -> bool:
    """Files that likely hold secrets: the constitution decides (Article 1.3)."""
    return constitution.is_sensitive_path(p)


def read_file(path: str, max_chars: int = 6000) -> str:
    """Read a text or code file (first 6000 characters by default)."""
    p = _resolve(path)
    if not p.is_file():
        return f"Error: file not found: {p}"
    if _is_sensitive(p):
        return f"Error: refusing to read {p.name}: it may contain secrets."
    with p.open("rb") as fh:
        head = fh.read(4096)
    if b"\x00" in head:
        return f"Error: {p.name} looks like a binary file. Use extract_pdf_text or read_docx_text for documents."
    text = p.read_text(encoding="utf-8", errors="replace")
    limit = int(max_chars)
    return text[:limit] + (f"\n...[truncated, {len(text) - limit} more characters]" if len(text) > limit else "")


def search_files(pattern: str, path: str = ".", max_results: int = 30) -> str:
    """Find lines containing text in the files under a folder (case-insensitive)."""
    root = _resolve(path)
    if not root.is_dir():
        return f"Error: directory not found: {root}"
    needle = str(pattern).lower()
    if not needle:
        return "Error: pattern is empty."
    hits, scanned = [], 0
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in IGNORED_DIRS]
        for name in filenames:
            fp = Path(dirpath) / name
            if _is_sensitive(fp) or fp.suffix.lower() not in TEXT_EXTENSIONS:
                continue
            scanned += 1
            if scanned > 3000:
                break
            try:
                if fp.stat().st_size > 1_000_000:
                    continue
                with fp.open(encoding="utf-8", errors="replace") as fh:
                    for lineno, line in enumerate(fh, 1):
                        if needle in line.lower():
                            hits.append(f"{fp.relative_to(root).as_posix()}:{lineno}: {line.strip()[:160]}")
                            if len(hits) >= int(max_results):
                                return "\n".join(hits) + f"\n...[stopped at {int(max_results)} matches]"
            except OSError:
                continue
    return "\n".join(hits) if hits else f"No matches for '{pattern}' in {scanned} files."


def get_system_info() -> str:
    """Report CPU, RAM, disk space and operating system of this computer."""
    import psutil

    vm = psutil.virtual_memory()
    disk = psutil.disk_usage(str(config.BASE_DIR))
    return "\n".join([
        f"OS: {platform.system()} {platform.release()} ({platform.version()})",
        f"CPU: {platform.processor() or 'unknown'}, {psutil.cpu_count(logical=False)} cores / {psutil.cpu_count()} threads, "
        f"{psutil.cpu_percent(interval=0.3):.0f}% busy",
        f"RAM: {vm.total / 2**30:.1f} GB total, {vm.available / 2**30:.1f} GB free ({vm.percent:.0f}% used)",
        f"Disk (project drive): {disk.total / 2**30:.0f} GB total, {disk.free / 2**30:.0f} GB free",
        f"Python: {platform.python_version()}",
    ])


def write_file(path: str, content: str) -> str:
    p = _resolve(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return f"Wrote {len(content)} chars to {p}"


def get_directory_tree(path: str = ".", max_depth: int = 3, max_entries: int = 200) -> str:
    """List the files and folders under path (ignores .git, node_modules, venvs...)."""
    root = _resolve(path)
    if not root.is_dir():
        return f"Error: directory not found: {root}"
    lines = [f"{root.name or root}/"]
    count = 0

    def walk(directory: Path, depth: int, prefix: str) -> None:
        nonlocal count
        if depth > int(max_depth):
            return
        try:
            entries = sorted(directory.iterdir(), key=lambda e: (e.is_file(), e.name.lower()))
        except OSError:
            return
        for entry in entries:
            if entry.name in IGNORED_DIRS:
                continue
            if count >= int(max_entries):
                return
            count += 1
            if entry.is_dir():
                lines.append(f"{prefix}{entry.name}/")
                walk(entry, depth + 1, prefix + "  ")
            else:
                lines.append(f"{prefix}{entry.name}")

    walk(root, 1, "  ")
    if count >= int(max_entries):
        lines.append("  ...[truncated]")
    return "\n".join(lines)


# --- Registry exposed to models ------------------------------------------

MODEL_TOOLS = {
    "web_search": web_search,
    "git_status": git_status,
    "git_pull": git_pull,
    "git_log": git_log,
    "git_diff": git_diff,
    "get_directory_tree": get_directory_tree,
    "run_ruff_check": run_ruff_check,
    "run_pytest": run_pytest,
    "extract_pdf_text": extract_pdf_text,
    "read_docx_text": read_docx_text,
    "read_file": read_file,
    "search_files": search_files,
    "fetch_webpage": fetch_webpage,
    "get_system_info": get_system_info,
}
assert "git_push" not in MODEL_TOOLS


def describe_tools(registry=None) -> str:
    """One line per tool for the prompt: name(signature) - first doc line."""
    lines = []
    for name, func in (registry or MODEL_TOOLS).items():
        sig = inspect.signature(func)
        params = ", ".join(
            p.name if p.default is inspect.Parameter.empty else f"{p.name}={p.default!r}"
            for p in sig.parameters.values()
        )
        doc = (inspect.getdoc(func) or "").splitlines()[0] if inspect.getdoc(func) else ""
        lines.append(f"- {name}({params}): {doc}")
    return "\n".join(lines)


def parse_tool_calls(text: str, limit: int = 3) -> list:
    """Every `TOOL_CALL: name(args)` line in model output, as [(name, args, kwargs)].

    Arguments are literals only, parsed with ast, never evaluated. Lines that are
    not a well-formed call are skipped.
    """
    calls = []
    for line in text.splitlines():
        line = line.strip()
        if not line.upper().startswith("TOOL_CALL:"):
            continue
        body = line.split(":", 1)[1].strip()
        try:
            node = ast.parse(body, mode="eval").body
            if not isinstance(node, ast.Call) or not isinstance(node.func, ast.Name):
                continue
            args = [ast.literal_eval(a) for a in node.args]
            kwargs = {k.arg: ast.literal_eval(k.value) for k in node.keywords if k.arg}
        except (SyntaxError, ValueError):
            continue
        calls.append((node.func.id, args, kwargs))
        if len(calls) >= limit:
            break
    return calls


def parse_tool_call(text: str):
    """The first tool call in the text as (name, args, kwargs), or None."""
    calls = parse_tool_calls(text, limit=1)
    return calls[0] if calls else None


# Names small models habitually make up -> the real tools, first one present in the registry wins.
TOOL_ALIASES = {
    "read_pdf_text": ("extract_pdf_text",), "read_pdf": ("extract_pdf_text",), "pdf_to_text": ("extract_pdf_text",),
    "extract_pdf": ("extract_pdf_text",), "extract_text_from_pdf": ("extract_pdf_text",),
    "read_pdf_file": ("extract_pdf_text",), "pdf_text": ("extract_pdf_text",),
    "read_docx": ("read_docx_text",), "extract_docx_text": ("read_docx_text",), "read_word_file": ("read_docx_text",),
    "read_text_file": ("read_workspace_file", "read_file"), "open_file": ("read_workspace_file", "read_file"),
    "list_files": ("list_workspace", "get_directory_tree"), "list_directory": ("get_directory_tree", "list_workspace"),
    "list_dir": ("get_directory_tree", "list_workspace"), "ls": ("list_workspace", "get_directory_tree"),
    "search": ("web_search",), "google": ("web_search",), "websearch": ("web_search",), "search_web": ("web_search",),
    "fetch_url": ("fetch_webpage",), "fetch_page": ("fetch_webpage",), "get_webpage": ("fetch_webpage",),
    "read_webpage": ("fetch_webpage",), "fetch_website": ("fetch_webpage",), "open_url": ("fetch_webpage",),
    "grep": ("search_files",), "find_in_files": ("search_files",), "search_in_files": ("search_files",),
    "system_info": ("get_system_info",), "get_system_information": ("get_system_info",),
    "install_package": ("pip_install",), "pip": ("pip_install",),
}


def resolve_tool_name(name: str, registry) -> str:
    """The registry key a (possibly misspelled) tool name refers to, or ''."""
    import difflib

    if name in registry:
        return name
    for candidate in TOOL_ALIASES.get(name.lower(), ()):
        if candidate in registry:
            return candidate
    close = difflib.get_close_matches(name.lower(), list(registry), n=1, cutoff=0.8)
    return close[0] if close else ""


def call_tool(name: str, args=(), kwargs=None, registry=None) -> str:
    """Execute a model-requested tool; always returns text, never raises."""
    registry = MODEL_TOOLS if registry is None else registry
    real = resolve_tool_name(name, registry)
    func = registry.get(real)
    if func is None:
        return f"Error: unknown tool '{name}'. Available: {', '.join(registry)}"
    try:
        result = str(func(*args, **(kwargs or {})))
    except TypeError as exc:
        return f"Error: bad arguments for {name}: {exc}"
    except Exception as exc:
        return f"Error: {type(exc).__name__}: {exc}"
    if len(result) > config.TOOL_RESULT_MAX_CHARS:
        result = result[: config.TOOL_RESULT_MAX_CHARS] + "\n...[truncated]"
    return result
